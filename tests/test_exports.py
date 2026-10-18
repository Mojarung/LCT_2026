"""Выгрузки для проектировщика: книга Excel и DXF только со слоями результата.

Green Leaders отдают ведомость и таблицу проверок в .xlsx (числа - числами), konstantikkov -
DXF только со слоями результата в координатах исходника, чтобы подключить его внешней ссылкой
к своему генплану (лёгкий файл даже для Олимпийской деревни).
"""

from __future__ import annotations

from dataclasses import replace
from typing import TYPE_CHECKING

import ezdxf
import openpyxl
from shapely.geometry import box
from test_writer_layout import BOOK, _mixed_plan, _write

from green.domain.norms import LawnKind
from green.domain.planting import Lawn
from green.infrastructure.cad.export_validation import check_written_plan
from green.infrastructure.reports.artifacts import write_result_layers, write_workbook

if TYPE_CHECKING:
    from pathlib import Path

    import pytest

    from green.domain.planting import Plan


def test_workbook_has_the_schedule_and_every_check_with_numbers(tmp_path: Path) -> None:
    plan = _mixed_plan()
    path = write_workbook(tmp_path / "plan.xlsx", plan, BOOK)
    book = openpyxl.load_workbook(path)
    assert book.sheetnames == ["Ведомость", "Проверки"]
    schedule = list(book["Ведомость"].iter_rows(values_only=True))
    counts = [row[3] for row in schedule[1:] if isinstance(row[3], int)]
    assert sum(counts) >= len(plan.placements)
    checks = list(book["Проверки"].iter_rows(values_only=True))
    head = checks[0]
    assert "measured_m" in head
    measured = [row[head.index("measured_m")] for row in checks[1:]]
    assert measured
    assert all(isinstance(value, (int, float)) for value in measured if value is not None)


def test_result_layers_file_keeps_only_green_layers_in_source_coordinates(tmp_path: Path) -> None:
    plan = _mixed_plan()
    result = _write(tmp_path, plan)
    light = write_result_layers(result, tmp_path / "result_layers.dxf")
    doc = ezdxf.readfile(light)
    layers = {e.dxf.layer for e in doc.modelspace()}
    assert layers
    assert all(layer.startswith("GREEN_") for layer in layers)
    source = ezdxf.readfile(result)
    inserts = sorted(
        (round(e.dxf.insert.x, 3), round(e.dxf.insert.y, 3))
        for e in source.modelspace().query("INSERT")
        if e.dxf.layer.startswith("GREEN_")
    )
    copied = sorted(
        (round(e.dxf.insert.x, 3), round(e.dxf.insert.y, 3))
        for e in doc.modelspace().query("INSERT")
    )
    assert copied == inserts


def test_layers_file_saved_under_the_pending_name_is_reused_not_reread(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Сценарий пишет результат во временный «.result.pending.dxf» и потом переименовывает;
    писатель кладёт рядом «.result.pending_layers.dxf». Профиль Камчатской 17.10.2026: файл не
    находился, и результат на сотни мегабайт читался заново (18 с из 210)."""
    plan = _mixed_plan()
    source = tmp_path / "source.dxf"
    doc = ezdxf.new("R2018")
    doc.modelspace().add_line((0, 0), (100, 0))
    doc.saveas(source)
    from green.infrastructure.cad.writer import EzdxfPlanWriter  # noqa: PLC0415

    pending = tmp_path / ".result.pending.dxf"
    EzdxfPlanWriter(text_font="DejaVuSans").write(source, plan, BOOK, pending, unit_m=1.0)
    result = tmp_path / "result.dxf"
    pending.replace(result)

    def no_reread(*_args, **_kwargs):  # noqa: ANN002, ANN003, ANN202
        raise AssertionError("result.dxf read again")

    monkeypatch.setattr(ezdxf, "readfile", no_reread)
    light = write_result_layers(result, tmp_path / "result_layers.dxf")
    assert light.exists()
    assert not (tmp_path / ".result.pending_layers.dxf").exists()  # служебный файл не остаётся


def test_rebuild_replaces_the_stale_result_layers_file(tmp_path: Path) -> None:
    """Ревью 17.10.2026: после правки свежий «.result.pending_layers.dxf» лежал сиротой, а
    отдавался result_layers.dxf исходного прогона."""
    stale = tmp_path / "result_layers.dxf"
    stale.write_text("old", encoding="utf-8")
    fresh = tmp_path / ".result.pending_layers.dxf"
    fresh.write_text("new", encoding="utf-8")
    result = tmp_path / "result.dxf"
    result.write_text("plan", encoding="utf-8")
    write_result_layers(result, stale)
    assert stale.read_text(encoding="utf-8") == "new"
    assert not fresh.exists()


def _with_flowerbed() -> Plan:
    bed = Lawn(
        "L00001",
        1,
        LawnKind.FLOWERBED,
        box(0, 0, 10, 6),
        ("R-LAWN-FLOWER-001", "R-LAWN-DRAW-001"),
        mix="солнечный цветник: злаки и многолетники",
        plants=(("Котовник Фассена", "Nepeta x faassenii", 135), ("Очиток видный", "x", 108)),
    )
    return replace(_mixed_plan(), lawns=(bed,))


def test_workbook_lists_flowerbed_perennials_by_bed(tmp_path: Path) -> None:
    path = write_workbook(tmp_path / "plan.xlsx", _with_flowerbed(), BOOK)
    book = openpyxl.load_workbook(path)

    assert book.sheetnames == ["Ведомость", "Цветники", "Проверки"]
    rows = list(book["Цветники"].iter_rows(values_only=True))
    assert rows[1] == (
        "Ц1",
        "солнечный цветник: злаки и многолетники",
        60,
        "Котовник Фассена",
        "Nepeta x faassenii",
        135,
    )
    assert sum(row[5] for row in rows[1:]) == 243


def test_flowerbed_goes_to_its_own_layer_and_passes_the_export_check(tmp_path: Path) -> None:
    plan = _with_flowerbed()
    result = _write(tmp_path, plan)
    doc = ezdxf.readfile(result)

    hatches = [e for e in doc.modelspace().query("HATCH") if e.dxf.layer == "GREEN_FLOWERBED"]
    assert len(hatches) == 1
    assert not [e for e in doc.modelspace().query("HATCH") if e.dxf.layer == "GREEN_LAWN"]
    texts = [e.dxf.text for e in doc.modelspace().query("TEXT") if e.dxf.layer == "GREEN_FLOWERBED"]
    assert texts == ["Ц1"]
    assert check_written_plan(result, plan, unit_m=1.0).ok


# Узоры штриховок LibreCAD 2.2 (/usr/share/librecad/patterns в образе docker/cadcheck):
# штриховку с узором не из списка LibreCAD выбрасывает при открытии («Invalid objects
# removed: 1» - цветник с узором STARS, 18.10.2026).
LIBRECAD_PATTERNS = {"ANSI31", "BRICK", "DOLMIT", "EARTH", "ESCHER", "GRASS", "HEX", "SAND"}


def test_result_hatches_use_patterns_librecad_knows(tmp_path: Path) -> None:
    doc = ezdxf.readfile(_write(tmp_path, _with_flowerbed()))

    patterns = {
        h.dxf.pattern_name.upper() for h in doc.modelspace().query("HATCH") if not h.dxf.solid_fill
    }
    assert patterns
    assert patterns <= LIBRECAD_PATTERNS
