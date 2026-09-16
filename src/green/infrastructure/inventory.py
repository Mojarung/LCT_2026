"""Перечётная ведомость: сколько каких деревьев уже растёт на участке.

Ведомость нужна квотам разнообразия: на Берзарина из 868 записей 304 - клён ясенелистный,
и без этого знания сервис спокойно добавил бы к ним ещё клёнов. Формат ведомостей пилота -
книга Excel с шапкой «№№ | Наименование | Кол-во в шт. | ... | Заключение»; шапка ищется по
слову «Наименование», потому что до неё идут титул и коэффициенты сметы.

Породы в ведомости записаны по-русски и часто одним словом («Вяз», «Клен»), поэтому
сопоставление идёт от точного названия к роду, а всё, что не опознано («Самосев до 8 см.»,
«поросль»), возвращается отдельным списком, а не молча теряется.
"""

from __future__ import annotations

import re
from collections import Counter
from typing import TYPE_CHECKING, Any

from green.application.errors import InputError
from green.application.ports import InventoryCounts

if TYPE_CHECKING:
    from collections.abc import Iterator, Sequence
    from pathlib import Path

    from green.domain.planting import Species

_HEADER = "наименование"
_COUNT = "кол-во"
_VERDICT = "заклю"
_REMOVE = re.compile(r"выруб|удал|снос")
_SORT = re.compile(r"[\"'«][^\"'»]*[\"'»]|\([^)]*\)")
_NOISE = re.compile(r"\bкуст\b|\bдер\b|\bшт\b|\bс\d+\b|крупномер")
_SUMMARY = re.compile(r"всего|итого|кв\.?\s*м|трав\.|покров|деревьев|кустарников|газон")
_UNIDENTIFIED = re.compile(r"самосев|поросл|пень|сухостой")
# В ведомости между таблицами идут строки-заголовки («Вырубить», «Сохранить») и строки
# других разделов - размеры ям «0,5х0,5х0,4», «5-ти м зона». Породой они не являются.
_NOT_SPECIES = re.compile(r"^\d|^выруб|^сохран|зона|^прим|^итог")
_MIN_NAME_LETTERS = 3
_MAX_HEADER_ROW = 40


def read_inventory(path: Path, catalog: Sequence[Species]) -> InventoryCounts:
    rows = list(_rows(path))
    if not rows:
        raise InputError(f"Перечётная ведомость {path.name} пуста или не читается")
    header = _header(rows)
    if header is None:
        raise InputError(
            f"В ведомости {path.name} не найдена строка заголовков со словом «Наименование»"
        )
    index, columns = header
    matched = Counter[str]()
    unmatched = Counter[str]()
    approximate: dict[str, str] = {}
    lookup = _lookup(catalog)
    read = removed = without_count = rows_matched = rows_unmatched = 0
    for row in rows[index + 1 :]:
        name = _text(row, columns.get("name"))
        folded = name.casefold()
        if not name or _SUMMARY.search(folded) or _NOT_SPECIES.search(folded):
            continue
        if len(re.findall(r"[а-яa-z]", folded)) < _MIN_NAME_LETTERS:
            continue
        read += 1
        if _REMOVE.search(_text(row, columns.get("verdict")).casefold()):
            removed += 1
            continue
        count = _count(_text(row, columns.get("count")))
        if count is None:
            without_count += 1
            count = 1
        code, exact = _match(name, lookup)
        if code is None:
            unmatched[name] += count
            rows_unmatched += 1
            continue
        matched[code] += count
        rows_matched += 1
        if not exact:
            approximate[name] = code
    return InventoryCounts(
        matched=dict(sorted(matched.items())),
        unmatched=dict(sorted(unmatched.items())),
        approximate=dict(sorted(approximate.items())),
        rows_read=read,
        rows_removed=removed,
        rows_matched=rows_matched,
        rows_unmatched=rows_unmatched,
        rows_without_count=without_count,
    )


def _rows(path: Path) -> Iterator[list[Any]]:
    suffix = path.suffix.casefold()
    # Чтение книг Excel нужно раз в прогон и только с --inventory, поэтому обе библиотеки
    # грузятся по факту обращения, а не при старте CLI.
    if suffix == ".xls":
        import xlrd  # noqa: PLC0415

        sheet = xlrd.open_workbook(path).sheet_by_index(0)
        for index in range(sheet.nrows):
            yield [sheet.cell_value(index, column) for column in range(sheet.ncols)]
    elif suffix in {".xlsx", ".xlsm"}:
        from openpyxl import load_workbook  # noqa: PLC0415

        worksheet = load_workbook(path, read_only=True, data_only=True).worksheets[0]
        for row in worksheet.iter_rows(values_only=True):
            yield list(row)
    else:
        raise InputError(f"Перечётка читается из .xls и .xlsx, получено {path.suffix}")


def _header(rows: Sequence[list[Any]]) -> tuple[int, dict[str, int]] | None:
    for index, row in enumerate(rows[:_MAX_HEADER_ROW]):
        cells = [str(cell or "").strip().casefold() for cell in row]
        if not any(_HEADER in cell for cell in cells):
            continue
        columns = {
            "name": next(i for i, cell in enumerate(cells) if _HEADER in cell),
            "count": next((i for i, cell in enumerate(cells) if _COUNT in cell), -1),
            "verdict": next((i for i, cell in enumerate(cells) if _VERDICT in cell), -1),
        }
        return index, {key: value for key, value in columns.items() if value >= 0}
    return None


def _text(row: Sequence[Any], column: int | None) -> str:
    if column is None or column >= len(row):
        return ""
    return str(row[column] or "").strip()


def _count(value: str) -> int | None:
    try:
        count = int(float(value.replace(",", ".")))
    except ValueError:
        return None
    return count if count > 0 else None


def _normalize(name: str) -> str:
    folded = _SORT.sub("", name).casefold().replace("ё", "е")
    cleaned = _NOISE.sub("", folded)
    return " ".join(re.findall(r"[а-яa-z]+", cleaned))


def _lookup(catalog: Sequence[Species]) -> tuple[dict[str, str], dict[str, str]]:
    """Два словаря: точный по названию и запасной по русскому роду (первому слову)."""
    exact: dict[str, str] = {}
    by_genus: dict[str, str] = {}
    for species in sorted(catalog, key=lambda s: (-s.pilot_streets, s.code)):
        normalized = _normalize(species.name_ru)
        if not normalized:
            continue
        exact.setdefault(normalized, species.code)
        exact.setdefault(" ".join(normalized.split()[:2]), species.code)
        by_genus.setdefault(normalized.split()[0], species.code)
    return exact, by_genus


def _match(name: str, lookup: tuple[dict[str, str], dict[str, str]]) -> tuple[str | None, bool]:
    exact, by_genus = lookup
    normalized = _normalize(name)
    if not normalized or _UNIDENTIFIED.search(normalized):
        return None, False
    words = normalized.split()
    for size in (2, 1):
        key = " ".join(words[:size])
        if key in exact:
            return exact[key], size == len(words) or size == 2  # noqa: PLR2004 - род и вид
    if words and words[0] in by_genus:
        return by_genus[words[0]], False
    return None, False
