"""Пояснительная записка прогона в DOCX: участок, нормативная база, решения, ведомость.

Проектировщик сдаёт в ДПиООС пояснительную записку (у пилота - «Раздел 1.1 ПЗ» Багрицкого и
Олимпийской деревни, записка 2-й Прядильной): общие сведения, нормативная база, принятые
решения, ведомость элементов озеленения (ГОСТ 21.508-2020, форма 9). Записка собирается из
тех же данных, что отчёт интерпретаций (interpretation_report.note_sections), без своих
расчётов: числа в записке и в report.html одни и те же.
"""

from __future__ import annotations

from typing import TYPE_CHECKING

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.shared import Cm, Pt

from green.infrastructure.reports.interpretation_report import note_sections

if TYPE_CHECKING:
    from pathlib import Path

    from docx.document import Document as DocxDocument

    from green.application.results import RunReport
    from green.infrastructure.reports.interpretation_report import NoteSection

_FONT = "Times New Roman"
_FONT_PT = 11
_TABLE_PT = 9


def write_note(path: Path, report: RunReport) -> Path:
    """Записка одним файлом .docx; таблицы - сеткой, шрифт - как у записок пилота."""
    document = Document()
    _page(document)
    document.core_properties.title = f"Пояснительная записка: {report.source_name}"
    document.core_properties.author = "green"
    document.add_heading(f"Пояснительная записка к плану озеленения: {report.source_name}", 0)
    for part in note_sections(report):
        document.add_heading(part.title, level=1)
        for text in part.text:
            document.add_paragraph(text)
        for item in part.items:
            document.add_paragraph(item, style="List Bullet")
        if part.head:
            _table(document, part)
    document.save(str(path))
    return path


def _page(document: DocxDocument) -> None:
    section = document.sections[0]
    # Альбомный лист: в ведомости и балансе по семь граф.
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    for margin in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, margin, Cm(2))
    style = document.styles["Normal"]
    style.font.name = _FONT
    style.font.size = Pt(_FONT_PT)


def _table(document: DocxDocument, part: NoteSection) -> None:
    table = document.add_table(rows=1, cols=len(part.head))
    table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells, part.head, strict=True):
        cell.text = text
    for row in part.rows:
        for cell, text in zip(table.add_row().cells, row, strict=True):
            cell.text = text
    for index, row in enumerate(table.rows):
        for cell in row.cells:
            for run in (r for paragraph in cell.paragraphs for r in paragraph.runs):
                run.font.size = Pt(_TABLE_PT)
                run.bold = index == 0


__all__ = ["write_note"]
