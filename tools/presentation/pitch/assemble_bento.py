# ruff: noqa: INP001, T201 - standalone deck assembler
"""Вторая версия презентации: обязательные слайды как в assemble.py, свободные - бенто-сетками.

Обязательные слайды 7-11 шаблона заполняет deck/ так же, как в assemble.py; свободные рисует
native_bento.py вместо native.py. PDF - тот же PPTX через LibreOffice.

Запуск (после shots.mjs и assemble.py: фоны bg/ и снимки shots/ - общие):
    uv run --with python-pptx --with pillow --with pymupdf python \
        tools/presentation/pitch/assemble_bento.py "ЛЦТ2026 Шаблон презентации.pptx"
Выход: out/presentation/pitch-bento/green-lct2026-bento.pptx и .pdf.
"""

from __future__ import annotations

import sys
from pathlib import Path

import pymupdf
from pptx import Presentation

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parent / "deck"))
sys.path.insert(0, str(HERE))

import assemble  # noqa: E402 - модули колоды лежат рядом, путь добавлен выше
import build_pptx as deck  # noqa: E402
import native_bento  # noqa: E402

OUT = assemble.ROOT / "out" / "presentation" / "pitch-bento"


def build(template: Path) -> Presentation:
    prs = Presentation(str(template))
    deck.drop_slides(prs, set(deck.MANDATORY))
    title_s, about_s, team_s, history_s, short_s = list(prs.slides)
    assets = {"plan_band": assemble.plan_band(), **deck.poser_logo()}
    deck.fill_title_slide(title_s, assets)
    deck.fill_about_slide(about_s, assets)
    deck.fill_team_slide(team_s, [deck.portrait(*spec) for spec in deck.TEAM_PHOTOS])
    deck.fill_history_slide(history_s)
    deck.fill_short_slide(short_s)
    native_bento.add_slides(prs)
    return prs


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    out_pptx = OUT / "green-lct2026-bento.pptx"
    prs = build(Path(sys.argv[1]))
    prs.core_properties.title = "green - озеленение улиц по нормам"
    prs.core_properties.author = "MISIS MOJARUNG"
    prs.save(str(out_pptx))
    out_pdf = assemble.to_pdf(out_pptx)
    with pymupdf.open(out_pdf) as doc:
        pages = doc.page_count
    print(f"{out_pptx}: {len(prs.slides)} слайдов, {out_pptx.stat().st_size / 1e6:.1f} МБ")
    print(f"{out_pdf}: {pages} страниц, {out_pdf.stat().st_size / 1e6:.1f} МБ")


if __name__ == "__main__":
    main()
