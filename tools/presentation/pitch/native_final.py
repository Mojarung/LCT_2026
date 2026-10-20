# ruff: noqa: INP001, E501, ISC004 - модуль колоды рядом с assemble_final.py; текст слайдов длинными строками
"""Свободные слайды финальной версии: по каждому слайду - native.py, native_bento.py или их смесь.

Выбор слайдов - по пометкам команды к двум прежним сборкам. Где слайд собран из частей обеих,
функция лежит здесь; где взят целиком - вызывается функция исходного модуля.
"""

from __future__ import annotations

from typing import TYPE_CHECKING

import native as m
import native_bento as b
from native import (
    BLUSH,
    BODY,
    IM_CODES,
    IM_ROWS,
    LAV,
    LEFT,
    MEDIUM,
    MIST,
    NIGHT,
    PINK,
    PLUM,
    REVIEW_CROP,
    RIGHT,
    ROSE,
    SEMI,
    VISUAL_CHIP,
    VISUAL_PANEL,
    VISUAL_SCENE,
    WHITE,
    WIDTH,
    X_BOLD,
    Crop,
    Style,
    circle_number,
    effect_stats,
    facts,
    fill_frame,
    notes,
    para,
    picture,
    rect,
    rgb,
    street_minutes,
    style_chart,
    text,
)
from native_bento import (
    ACCENT,
    DEEP,
    GLASS,
    LIGHT,
    MARGIN,
    SLIDE_W,
    caption,
    card,
    cell,
    inset,
    note,
    shot,
    stat,
)
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

if TYPE_CHECKING:
    from collections.abc import Callable

    from pptx.presentation import Presentation as Deck


# --------------------------------------------------------------------------- 01 описание


def s_steps(prs: Deck, f: dict[str, str]) -> None:
    """Кружки шагов основной версии, текст шагов на карточках, внизу плитка бенто."""
    s = m.new_slide(prs, 1, "Семь шагов до решения")
    steps = (
        (
            "Точный парсинг",
            "Учёт 100% примитивов чертежа и сквозная сверка геометрии с исходным DXF.",
        ),
        (
            "Семантика слоёв",
            f"Распознавание сетей, бортов, застройки и {f['symbols']} типов условных знаков геоподосновы.",
        ),
        ("Анализ покрытий", "Газоны, тротуары, проезды и зоны с открытым грунтом под посадки."),
        ("Сценарии посадки", "Аллея вдоль бортов, куртины на газонах, живые изгороди, цветники."),
        (
            "Подбор пород (MILP)",
            f"Подбор из {f['species']} видов с учётом инсоляции, габаритов и городских квот.",
        ),
        (
            "Аудит и валидация",
            f"Финальная проверка плана по {f['rules_total']} правилам: {f['violations']} нарушений на выходе.",
        ),
        (
            "Экспорт и отчётность",
            "Слои GREEN_*, ведомость, CSV, записка DOCX и протокол со ссылками на акты.",
        ),
    )
    gap = 0.14
    w = (WIDTH - gap * 6) / 7
    for i, (head, sub) in enumerate(steps):
        x = LEFT + i * (w + gap)
        circle_number(s, x + (w - 0.56) / 2, 1.9, 0.56, str(i + 1), size=16)
        if i < len(steps) - 1:
            dash = s.shapes.add_connector(
                MSO_CONNECTOR.STRAIGHT,
                Inches(x + (w + 0.56) / 2 + 0.08),
                Inches(2.18),
                Inches(x + w + gap + (w - 0.56) / 2 - 0.08),
                Inches(2.18),
            )
            dash.line.color.rgb = rgb(BLUSH)
            dash.line.width = Pt(2)
        card(s, (x, 2.66, w, 2.46), LIGHT)
        text(s, (x + 0.1, 2.76, w - 0.2, 0.62), [para(head, Style(X_BOLD, 12.5, NIGHT))], line=1.0)
        text(s, (x + 0.1, 3.42, w - 0.2, 1.65), [para(sub, Style(MEDIUM, 10.5, BODY))], line=1.2)
    _, by, _, bh = cell(0, 12, 4, 2)
    box = (LEFT, by, WIDTH, bh)
    card(s, box, DEEP)
    x, y, w, h = inset(box, 0.3)
    text(
        s,
        (x, y, 4.3, h),
        [
            para("Математическая оптимизация", Style(X_BOLD, 18, WHITE)),
            para("вместо машинного обучения", Style(X_BOLD, 18, ROSE)),
        ],
        line=1.05,
        anchor=MSO_ANCHOR.MIDDLE,
    )
    text(
        s,
        (x + 4.6, y, w - 4.6, h),
        [
            para(
                "В принятом плане Берзарина нормы отступов нарушены в 391 посадке из 554: обучать "
                "нейросеть на таких данных нельзя. Солвер MILP (HiGHS) строго соблюдает нормативные "
                "пороги, каждое дерево проверяемо. Локальная языковая модель только подсказывает "
                "класс незнакомого слоя.",
                Style(MEDIUM, 12, MIST),
            )
        ],
        line=1.2,
        anchor=MSO_ANCHOR.MIDDLE,
    )
    notes(
        s, "Семь шагов конвейера. Ключевое: генерация и проверка разделены, проверка независимая."
    )


def s_explain(prs: Deck, f: dict[str, str]) -> None:
    """Бенто-версия, главное число - на тёмной карточке вместо розовой."""
    s = b.new_slide(prs, 1, "Проверка норм для каждого растения")
    shot(
        s,
        "explain-map.png",
        cell(0, 4, 0, 6),
        Crop((0.3, 0.5), 1.3, (0.5, 0.5)),
        "Выноска: замер до сети",
    )
    shot(s, "explain-card.png", cell(4, 3, 0, 6), Crop((0.5, 0.0)))
    note(
        s,
        cell(7, 5, 0, 2),
        "Наглядные замеры на плане и в карточке",
        "Сервис рассчитывает расстояние до сетей (теплосеть: факт 3,47 м ≥ норма 2,50 м) и прикрепляет пункт документа с дословной цитатой.",
        eyebrow="посадка",
        head_size=12.5,
        body_size=10.5,
    )
    note(
        s,
        cell(7, 5, 2, 2),
        "Аргументированное выявление нарушений",
        "Отказ со ссылкой на регламент: «до силового кабеля 0,51 м при норме 2,00 м» (R-POWER-TREE-001, СП 42.13330.2016, п. 9.6, табл. 9.1).",
        eyebrow="отказ",
        head_size=12.5,
        body_size=10.5,
    )
    stat(
        s,
        cell(7, 5, 4, 2),
        f"{f['rules_total']} правил · {f['acts']} актов",
        f"Без «галлюцинаций»: в базе {f['rules_total']} правил из {f['acts']} нормативных актов. Формулировки собирает строгий алгоритм, а не нейросеть: результат детерминирован и проверяем.",
        kind=DEEP,
        size=22,
        label_size=10,
    )
    notes(s, "Интерпретируемость: у каждой посадки и отказа замер, правило, акт, пункт, цитата.")


# --------------------------------------------------------------------------- 02 маркетинг


def s_users(prs: Deck, f: dict[str, str]) -> None:
    """Роли основной версии, ответы заказчика - на карточках крупным текстом."""
    s = m.new_slide(prs, 2, "Кому нужен сервис")
    roles = (
        (
            "Проектировщик благоустройства",
            "Загружает чертёж, автоматически получает схему посадок, вносит правки "
            "в веб-интерфейсе и возвращает результат в привычный CAD.",
            f"Скорость: {f['minutes_median']} мин на улицу",
        ),
        (
            "Специалист ДПиООС",
            "Автоматически проверяет отступы от сетей и нормативную базу. За минуты "
            "проверяет готовые чертежи подрядчиков перед согласованием.",
            "Контроль: аудит каждой посадки",
        ),
        (
            "Согласующие ведомства",
            "Оценивают баланс и объёмы озеленения, ведомость, записку DOCX и будущий облик "
            "улицы в интерактивном 3D-виде.",
            "Наглядность: ведомость и 3D",
        ),
    )
    gap = 0.3
    w = (WIDTH - 2 * gap) / 3
    for i, (head, body, metric) in enumerate(roles):
        x = LEFT + i * (w + gap)
        rect(s, (x, 1.9, w, 2.65), WHITE, line=(BLUSH, 1.25), radius=0.18)
        circle_number(s, x + 0.3, 2.12, 0.5, str(i + 1), size=14)
        text(
            s,
            (x + 0.95, 2.12, w - 1.2, 0.55),
            [para(head, Style(X_BOLD, 14, NIGHT))],
            anchor=MSO_ANCHOR.MIDDLE,
            line=1.0,
        )
        text(s, (x + 0.3, 2.8, w - 0.6, 1.3), [para(body, Style(MEDIUM, 11.5, BODY))], line=1.22)
        text(s, (x + 0.3, 4.1, w - 0.6, 0.35), [para(metric, Style(SEMI, 12, PINK))])
    text(s, (LEFT, 4.68, WIDTH, 0.35), [para("Валидация болей заказчика", Style(SEMI, 13, PLUM))])
    quotes = (
        (
            "Топ-1 и Топ-2",
            "главные причины возврата проектов: нарушение отступов от сетей и ошибки в объёмах посадок",
        ),
        (
            "0",
            "свободная ниша: у заказчика не было попыток автоматизировать проектирование озеленения",
        ),
        ("DXF", "работа в стандартном AutoCAD без плагинов: результат - в обычных слоях чертежа"),
        (
            "до 60 мин",
            f"ориентир на генерацию плана улицы, до 30 мин на подбор мест; сервис - {f['minutes_median']} мин",
        ),
    )
    qgap = 0.2
    qw = (WIDTH - 3 * qgap) / 4
    for i, (big, label) in enumerate(quotes):
        x = LEFT + i * (qw + qgap)
        box = (x, 5.08, qw, 1.82)
        card(s, box, "soft")
        bx, by, bw, bh = inset(box, 0.18)
        text(
            s,
            (bx, by, bw, bh),
            [para(big, Style(X_BOLD, 20, PINK)), para(label, Style(MEDIUM, 11, NIGHT))],
            line=1.12,
            after=4,
        )
    notes(
        s,
        "Три пользователя из ТЗ, у каждого свой вход. Ниже - ответы заказчика (docs/notes/15): "
        "главная боль - отступы, автоматизации не было, нужен обычный DXF.",
    )


def s_visual(prs: Deck, f: dict[str, str]) -> None:  # noqa: ARG001 - общая подпись слайдов
    """Бенто-сетка: свежее фото ИИ, кадр 3D к нему, весна и лето; без кадров возраста."""
    s = b.new_slide(prs, 2, "3D-визуализация и фотореализм", bg="dark")
    w = 0.45 + len(VISUAL_CHIP) * 0.098
    chip = rect(s, (SLIDE_W - MARGIN - w, 1.12, w, 0.4), LAV, radius=0.2)
    fill_frame(
        chip.text_frame,
        [para(VISUAL_CHIP, Style(SEMI, 10.5, WHITE))],
        align=PP_ALIGN.CENTER,
        anchor=MSO_ANCHOR.MIDDLE,
    )
    chip.text_frame.word_wrap = False
    shot(
        s,
        "fresh/fresh-photo-1.png",
        cell(0, 6, 0, 4),
        Crop(),
        "фото ИИ по кадру 3D-вида",
        dark=True,
    )
    shot(
        s,
        "extra/scene-city.png",
        cell(6, 3, 0, 2),
        Crop(scale=1.3, origin=(0.5, 0.2)),
        "город вокруг: OpenStreetMap",
        dark=True,
    )
    shot(
        s,
        "extra/scene-panel.png",
        cell(9, 3, 0, 2),
        Crop(scale=1.7, origin=(1.0, 0.75)),
        "пульт сцены",
        dark=True,
    )
    box = cell(6, 6, 2, 2)
    card(s, box, GLASS)
    text(
        s,
        inset(box, 0.2),
        [
            [(VISUAL_PANEL[0], Style(SEMI, 10, WHITE)), (VISUAL_PANEL[1], Style(MEDIUM, 10, MIST))],
            para(VISUAL_SCENE, Style(MEDIUM, 10, MIST)),
        ],
        line=1.08,
        after=3,
        anchor=MSO_ANCHOR.MIDDLE,
    )
    shot(
        s, "fresh/fresh-3d-1.png", cell(0, 4, 4, 2), Crop(), "кадр 3D-вида к этому фото", dark=True
    )
    shot(
        s,
        "extra/season-spring.png",
        cell(4, 4, 4, 2),
        m.tile_crop("extra/season-spring.png"),
        "весна",
        dark=True,
    )
    shot(
        s,
        "extra/season-summer.png",
        cell(8, 4, 4, 2),
        m.tile_crop("extra/season-summer.png"),
        "лето",
        dark=True,
    )
    notes(s, m.VISUAL_NOTES)


# --------------------------------------------------------------------------- 03 бизнес


def s_effect(prs: Deck, f: dict[str, str]) -> None:
    """Основная версия; числа эффекта - на полупрозрачных карточках."""
    s = m.new_slide(prs, 3, "Экологический эффект: тень и защита от пыли", bg="dark")
    top = sorted(m.street_curbs(), key=lambda r: r[2] - r[1], reverse=True)[:7]
    text(
        s,
        (LEFT, 1.9, 7.0, 0.6),
        [
            [
                ("Борта под кронами и кустарником, %: ", Style(SEMI, 13, WHITE)),
                ("было", Style(SEMI, 13, LAV)),
                (" и ", Style(SEMI, 13, WHITE)),
                ("с планом сервиса", Style(SEMI, 13, ROSE)),
            ],
            para("7 улиц с наибольшим приростом", Style(MEDIUM, 11, MIST)),
        ],
    )
    m.hbar_chart(
        s,
        (LEFT, 2.55, 7.0, 4.25),
        [r[0] for r in top],
        [("было", [r[1] for r in top], LAV), ("с планом сервиса", [r[2] for r in top], PINK)],
        dark=True,
        number_format="0.0",
        size=10.5,
        legend=False,
    )
    x, w = 8.0, RIGHT - 8.0
    for i, (big, label) in enumerate(effect_stats(f)):
        y = 1.88 + i * 1.22
        card(s, (x, y, w, 1.1), GLASS)
        text(s, (x + 0.22, y + 0.1, w - 0.44, 0.45), [para(big, Style(X_BOLD, 20, ROSE))])
        text(
            s,
            (x + 0.22, y + 0.52, w - 0.44, 0.55),
            [para(label, Style(MEDIUM, 10.5, WHITE))],
            line=1.12,
        )
    notes(
        s,
        "Эффект для города: тень, пылезащита, ярусность - то, что заказчик назвал признаками хорошего плана.",
    )


# --------------------------------------------------------------------------- 04 техника


def s_speed(prs: Deck, f: dict[str, str]) -> None:
    """Только график минут по улицам; ориентир 60 мин - пунктиром, 30 мин не рисуется."""
    s = m.new_slide(prs, 4, "Результаты пилота на реальных данных")
    rows = sorted(street_minutes(f), key=lambda r: r[1])
    text(
        s,
        (LEFT, 1.9, WIDTH, 0.6),
        [
            para(
                f"Минуты прогона по {len(rows)} улицам пилота, медиана {f['minutes_median']}",
                Style(SEMI, 14, NIGHT),
            ),
            para(
                "Пунктир - ориентир заказчика: 60 мин на генерацию плана улицы",
                Style(MEDIUM, 11.5, BODY),
            ),
        ],
    )
    data = CategoryChartData(number_format="0.0")
    data.categories = [r[0] for r in rows]
    data.add_series("минуты", [r[1] for r in rows])
    chart = s.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(LEFT), Inches(2.5), Inches(WIDTH), Inches(4.45), data
    ).chart
    style_chart(chart, dark=False, size=10)
    chart.has_legend = False
    plot = chart.plots[0]
    plot.gap_width = 45
    plot.has_data_labels = True
    labels = plot.data_labels
    labels.number_format = "0.0"
    labels.number_format_is_linked = False
    labels.position = XL_LABEL_POSITION.OUTSIDE_END
    labels.font.size = Pt(10)
    labels.font.bold = True
    labels.font.color.rgb = rgb(NIGHT)
    m.paint(plot.series[0], PINK)
    val = chart.value_axis
    val.minimum_scale = 0
    val.maximum_scale = 60
    val.major_unit = 60
    val.has_major_gridlines = True
    grid = val.major_gridlines.format.line
    grid.color.rgb = rgb(PLUM)
    grid.width = Pt(1.25)
    grid.dash_style = MSO_LINE_DASH_STYLE.DASH
    val.format.line.fill.background()
    val.tick_labels.font.size = Pt(10)
    val.tick_labels.number_format = "0"
    val.tick_labels.number_format_is_linked = False
    cat = chart.category_axis
    cat.format.line.color.rgb = rgb(BLUSH)
    cat.tick_labels.font.size = Pt(10)
    body_pr = cat._element.find(qn("c:txPr")).find(qn("a:bodyPr"))  # noqa: SLF001 - поворот подписей
    body_pr.set("rot", "-2700000")
    body_pr.set("vert", "horz")
    notes(
        s,
        "Время по всем улицам пилота, строгий профиль. Каждая улица - один прогон без ручных "
        "правок. Фрунзенская набережная в материалах пилота есть только листами PDF: "
        "tools/pdf_to_dxf.py собрал из пяти листов чертёж в метрах.",
    )


INPUT_ITEMS: tuple[tuple[str, str], ...] = (
    (
        "Файлы жюри - автоматически",
        "6 чертежей жюри (19-20 ПОДД, 236285, 236287, ИС ЗИЛ, два СПОЗУ) обработаны без ручных "
        "правок. Все проверки пройдены: исходник цел, план проверен, DXF совпадает с планом.",
    ),
    (
        "Классификаторы Москвы для ИМ, 2024",
        f"{IM_ROWS} строк, у {IM_CODES} кодов есть класс сервиса: слой с кодом «ТС 22 00 01» "
        "в имени читается как деревья.",
    ),
    (
        "Незнакомый слой",
        "Класс - по словам имени и по геометрии; спорное сервис показывает в таблице "
        "уточнения, решает человек.",
    ),
    (
        "Слои ГИС и кадастр",
        "GeoJSON и SHP: тип территории, охранные зоны газа и ЛЭП, границы участков.",
    ),
)


def s_inputs(prs: Deck, f: dict[str, str]) -> None:  # noqa: ARG001 - общая подпись слайдов
    """Основная версия с новым текстом: файлы жюри проходят без ручных правок."""
    s = m.new_slide(prs, 4, "Чужой чертёж и данные города")
    box = (LEFT, 1.95, 6.3, 4.85)
    picture(s, "review.png", box, REVIEW_CROP)
    caption(s, box, "уточнение объектов чертежа: класс каждого слоя и основание")
    x, w = 7.35, RIGHT - 7.35
    paras = []
    for head, body in INPUT_ITEMS:
        paras += [para(head, Style(X_BOLD, 13, NIGHT)), para(body, Style(MEDIUM, 11.5, BODY))]
    tb = text(s, (x, 1.95, w, 4.9), paras, line=1.15, after=1)
    for i, p in enumerate(tb.text_frame.paragraphs):
        if i and i % 2 == 0:  # отбивка перед заголовком пункта
            p.space_before = Pt(12)
    notes(
        s,
        "Файлы жюри: 19-20 ПОДД, 236285, 236287, ИС ЗИЛ и два СПОЗУ - каждый прогон без ручных "
        "правок, verify.json: исходные слои целы, план прошёл независимую проверку, DXF совпадает "
        "с планом. " + m.INPUT_NOTES,
    )


def s_workspace(prs: Deck, f: dict[str, str]) -> None:
    """Шесть экранов рабочего места ровной сеткой 3 x 2, подписи плашками."""
    s = b.new_slide(prs, 4, "Рабочее место в браузере")
    home, work, *small = m.workspace_tiles(f)
    tiles = (work, home, *small)
    for i, (name, crop, label) in enumerate(tiles):
        box = cell((i % 3) * 4, 4, (i // 3) * 3, 3)
        picture(s, name, box, crop, radius=b.R, shadow=False, border=BLUSH)
        caption(s, box, label)
    notes(
        s,
        "Главная: улица пилота или свой DXF/DWG и реестр прогонов. План: клик по кроне - замеры "
        "до сетей и пункт нормы. Вкладка «Эффект» - баланс было - стало и индекс качества. Окно "
        "«Отчёты и файлы прогона»: нормы CSV, отчёт по посадкам, записка DOCX, распознавание "
        "объектов JSON. «Как собран план» - какие варианты отброшены и почему. Каталог моделей "
        "растений - отдельная страница.",
    )


# --------------------------------------------------------------------------- 05 уникальность


def s_unique(prs: Deck, f: dict[str, str]) -> None:
    """Бенто-версия без повторов: главная карточка - документы в форме заказчика."""
    s = b.new_slide(prs, 5, "Чем green отличается", bg="dark")
    box = cell(0, 5, 0, 6)
    card(s, box, ACCENT)
    x, y, w, h = inset(box, 0.26)
    text(
        s,
        (x, y, w, 1.95),
        [
            para("ДОКУМЕНТЫ В ФОРМЕ ЗАКАЗЧИКА", Style(SEMI, 9.5, WHITE)),
            para("Записка DOCX и ведомость", Style(X_BOLD, 19, WHITE)),
            para(
                "Ведомость элементов озеленения по форме 9 ГОСТ 21.508-2020, пояснительная "
                "записка DOCX с пунктами норм по каждой посадке и книга Excel со всеми проверками.",
                Style(MEDIUM, 11, WHITE),
            ),
        ],
        line=1.1,
        after=4,
    )
    doc_box = (x, y + 2.0, w, h - 2.0)
    picture(
        s,
        "fresh/note-page1.png",
        doc_box,
        Crop((0.0, 0.0), 1.12, (0.0, 0.0)),
        radius=0.12,
        shadow=False,
    )
    caption(s, doc_box, "note.docx, первая страница")
    small = (
        (
            cell(5, 4, 0, 2),
            f"{f['entities_changed']} из {f['entities']}",
            "Исходник не меняется",
            "Новое пишется только в слои GREEN_*; каждый примитив сверяется по хэшу до и после записи.",
        ),
        (
            cell(9, 3, 0, 2),
            "Нормы в YAML",
            "Без релиза",
            "Новая редакция норм действует без пересборки Docker-образа.",
        ),
        (
            cell(5, 3, 2, 2),
            "СП 42: 2016 и 2026",
            "Две редакции",
            "Правила обеих редакций, переключатель в профиле.",
        ),
        (
            cell(8, 4, 2, 2),
            "Правка на лету",
            "Перенос с проверкой норм",
            "Перенесённую посадку и соседей в 30 м сервис перепроверяет, пересборка переписывает DXF.",
        ),
        (
            cell(5, 7, 4, 2),
            "Прикорневые барьеры",
            "Режим по СП 42, табл. 9.1, прим. 5 и 7",
            "Без барьера место у кабеля уходит в отказ с объяснением, при каком барьере оно допустимо; с барьером - посадка со ссылкой на примечание.",
        ),
    )
    for bx, metric, head, body in small:
        card(s, bx, GLASS)
        tx, ty, tw, th = inset(bx, 0.2)
        text(
            s,
            (tx, ty, tw, th),
            [
                para(metric, Style(X_BOLD, 16, ROSE)),
                para(head, Style(SEMI, 12, WHITE)),
                para(body, Style(MEDIUM, 10, MIST)),
            ],
            line=1.06,
            after=3,
        )
    notes(
        s,
        "Отличия, которых нет на других слайдах: документы в форме заказчика (записка DOCX, "
        "ведомость по форме 9 ГОСТ 21.508-2020, книга Excel), неизменность исходника, нормы "
        "данными, две редакции СП 42, правка с перепроверкой, режим прикорневых барьеров.",
    )


OUR_SLIDES: tuple[Callable[[Deck, dict[str, str]], None], ...] = (
    b.s_problem,  # 6
    m.s_solution,  # 7
    s_steps,  # 8
    s_explain,  # 9
    s_users,  # 10
    m.s_audit,  # 11
    s_visual,  # 12
    m.s_city,  # 13
    s_effect,  # 14
    m.s_deploy,  # 15
    s_speed,  # 16
    s_inputs,  # 17
    s_workspace,  # 18
    s_unique,  # 19
    b.s_roadmap,  # 20
)


def add_slides(prs: Deck) -> None:
    f = facts()
    for build in OUR_SLIDES:
        build(prs, f)
