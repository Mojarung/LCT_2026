# ruff: noqa: INP001, T201 - standalone deck assembler
"""Финальная версия презентации: обязательные слайды как в assemble.py, свободные - native_final.py.

native_final.py берёт каждый свободный слайд из native.py, native_bento.py или собирает из обоих.
PDF - тот же PPTX через LibreOffice; копия PDF кладётся в docs/ для формы сдачи.

Запуск (после shots.mjs и assemble.py: фоны bg/ и снимки shots/ - общие):
    uv run --with python-pptx --with pillow --with pymupdf python \
        tools/presentation/pitch/assemble_final.py "ЛЦТ2026 Шаблон презентации.pptx"
Выход: out/presentation/pitch-final/green-lct2026-final.pptx и .pdf,
docs/green-lct2026-final.pdf и docs/green-lct2026.pdf.
"""

from __future__ import annotations

import shutil
import sys
from pathlib import Path

import pymupdf
from pptx import Presentation

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parent / "deck"))
sys.path.insert(0, str(HERE))

import assemble  # noqa: E402 - модули колоды лежат рядом, путь добавлен выше
import build_pptx as deck  # noqa: E402
import native_final  # noqa: E402

OUT = assemble.ROOT / "out" / "presentation" / "pitch-final"
DOCS = assemble.ROOT / "docs"


def build(template: Path) -> Presentation:
    prs = Presentation(str(template))
    deck.drop_slides(prs, set(deck.MANDATORY))
    title_s, about_s, team_s, history_s, short_s = list(prs.slides)
    assets = {"plan_band": assemble.plan_band(), **deck.poser_logo()}
    deck.fill_title_slide(title_s, assets)
    deck.fill_about_slide(about_s, assets)
    deck.fill_team_slide(team_s, [deck.portrait(*spec) for spec in deck.TEAM_PHOTOS])
    deck.fill_history_slide(history_s)
    deck.fill_short_slide(short_s)
    native_final.add_slides(prs)
    return prs


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    out_pptx = OUT / "green-lct2026-final.pptx"
    prs = build(Path(sys.argv[1]))
    prs.core_properties.title = "green - озеленение улиц по нормам"
    prs.core_properties.author = "MISIS MOJARUNG"
    prs.save(str(out_pptx))
    out_pdf = assemble.to_pdf(out_pptx)
    with pymupdf.open(out_pdf) as doc:
        pages = doc.page_count
    for name in ("green-lct2026-final.pdf", "green-lct2026.pdf"):
        shutil.copyfile(out_pdf, DOCS / name)
    print(f"{out_pptx}: {len(prs.slides)} слайдов, {out_pptx.stat().st_size / 1e6:.1f} МБ")
    print(f"{out_pdf}: {pages} страниц, {out_pdf.stat().st_size / 1e6:.1f} МБ")


if __name__ == "__main__":
    main()
