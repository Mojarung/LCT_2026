# ruff: noqa: INP001, T201, SLF001 - standalone deck assembler; python-pptx has no API for slide order
"""Собирает презентацию к сдаче в шаблоне ЛЦТ 2026, целиком редактируемую в PowerPoint.

Слайды 7-11 шаблона обязательные (deck/mandatory.py): оформление шаблона, меняется только текст.
За ними слайд шаблона с рекомендуемой структурой и свободные слайды native.py по шести
разделам этой структуры: текст, фигуры, таблицы и графики PowerPoint, снимки сервиса - картинками.
PDF - тот же PPTX через LibreOffice.

Запуск (после shots.mjs; фоны bg/ - из шаблона):
    uv run --with python-pptx --with pillow --with pymupdf python \
        tools/presentation/pitch/assemble.py "ЛЦТ2026 Шаблон презентации.pptx"
Выход: out/presentation/pitch/green-lct2026.pptx и green-lct2026.pdf.
"""

from __future__ import annotations

import subprocess
import sys
from pathlib import Path

import pymupdf
from PIL import Image
from pptx import Presentation

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parent / "deck"))
sys.path.insert(0, str(HERE))

import build_pptx as deck  # noqa: E402 - модули колоды лежат рядом, путь добавлен выше
import native  # noqa: E402

ROOT = HERE.parents[2]
PITCH = ROOT / "out" / "presentation" / "pitch"
SOFFICE = Path(r"C:\Program Files\LibreOffice\program\soffice.exe")
STRUCTURE = 4  # слайд шаблона «Рекомендуемая структура презентации»


def plan_band() -> Path:
    """Полоса плана для слайда «О команде»: свежий снимок целой улицы, центральная часть."""
    deck.ASSETS.mkdir(parents=True, exist_ok=True)
    shot = Image.open(PITCH / "shots" / "plan-street-light.png").convert("RGB")
    return deck.jpg(shot.crop((1080, 380, 2220, 1420)), "plan_band.jpg")


def build(template: Path) -> Presentation:
    prs = Presentation(str(template))
    deck.drop_slides(prs, {STRUCTURE, *deck.MANDATORY})
    structure_s, title_s, about_s, team_s, history_s, short_s = list(prs.slides)
    ids = prs.slides._sldIdLst
    first = ids[0]
    ids.remove(first)
    ids.insert(len(deck.MANDATORY), first)  # структура - сразу после обязательных
    assets = {"plan_band": plan_band(), **deck.poser_logo()}
    deck.fill_title_slide(title_s, assets)
    deck.fill_about_slide(about_s, assets)
    deck.fill_team_slide(team_s, [deck.portrait(*spec) for spec in deck.TEAM_PHOTOS])
    deck.fill_history_slide(history_s)
    deck.fill_short_slide(short_s)
    native.fill_structure_slide(structure_s)
    native.add_slides(prs)
    return prs


def to_pdf(pptx: Path) -> Path:
    subprocess.run(  # noqa: S603 - LibreOffice по фиксированному пути, аргументы - наши файлы
        [
            str(SOFFICE),
            "--headless",
            "--convert-to",
            "pdf",
            str(pptx),
            "--outdir",
            str(pptx.parent),
        ],
        check=True,
        capture_output=True,
    )
    return pptx.with_suffix(".pdf")


def main() -> None:
    PITCH.mkdir(parents=True, exist_ok=True)
    out_pptx = PITCH / "green-lct2026.pptx"
    prs = build(Path(sys.argv[1]))
    prs.core_properties.title = "green - озеленение улиц по нормам"
    prs.core_properties.author = "MISIS MOJARUNG"
    prs.save(str(out_pptx))
    out_pdf = to_pdf(out_pptx)
    with pymupdf.open(out_pdf) as doc:
        pages = doc.page_count
    print(f"{out_pptx}: {len(prs.slides)} слайдов, {out_pptx.stat().st_size / 1e6:.1f} МБ")
    print(f"{out_pdf}: {pages} страниц, {out_pdf.stat().st_size / 1e6:.1f} МБ")


if __name__ == "__main__":
    main()
