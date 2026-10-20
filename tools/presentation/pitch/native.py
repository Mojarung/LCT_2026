# ruff: noqa: INP001, SLF001, PLR0913, E501, ISC004, PLR2004 - модуль колоды рядом с assemble.py; текст
# слайдов длинными строками, файл больше 800 строк из-за него же; у python-pptx нет API для фона,
# скругления картинок, прозрачности заливки и полей номера слайда.
"""Свободные слайды презентации нативными объектами PowerPoint: текст, фигуры, таблицы, графики.

Порядок - по рекомендуемой структуре шаблона ЛЦТ 2026: 01 описание, 02 маркетинг, 03 бизнес,
04 техника, 05 уникальность, 06 развитие. Картинками идут только снимки сервиса; их обрезка
задана свойствами кадра PowerPoint, и её можно поменять в самом PowerPoint.

Числа: facts.json и приложение B документации (прогон 20 улиц пилота), правила, акты и виды -
заново из config/, нормоконтроль Берзарина - docs/notes/24-audit.md, ответы заказчика -
docs/notes/15-organizers-qa.md.
"""

from __future__ import annotations

import json
from dataclasses import dataclass
from pathlib import Path
from typing import TYPE_CHECKING

from lxml import etree
from PIL import Image
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls, qn
from pptx.util import Inches, Pt

if TYPE_CHECKING:
    from collections.abc import Callable, Sequence

    from pptx.chart.chart import Chart
    from pptx.presentation import Presentation as Deck
    from pptx.shapes.autoshape import Shape
    from pptx.slide import Slide

ROOT = Path(__file__).resolve().parents[3]
PITCH = ROOT / "out" / "presentation" / "pitch"
SHOTS = PITCH / "shots"
BG = PITCH / "bg"
ASSETS = PITCH / "assets"
GENERATED = ROOT / "docs" / "documentation" / "generated"

# палитра шаблона ЛЦТ 2026 (слайд «Оформление»)
PINK = "FF0053"
ROSE = "FC3777"
BLUSH = "FFD6E4"
LAV = "8A83D1"
PLUM = "520978"
NIGHT = "310F53"
INK = "1C1D22"
BODY = "3C3548"
MUTE = "5E5670"
WHITE = "FFFFFF"
MIST = "E6DDF0"  # мелкий текст на фиолетовом фоне

X_BOLD = "Montserrat ExtraBold"
SEMI = "Montserrat SemiBold"
MEDIUM = "Montserrat Medium"
SANS = "Montserrat"
MONO = "Consolas"

LEFT = 0.67  # поле слайда, дюймы: 96 px сетки 1920
RIGHT = 12.67
WIDTH = RIGHT - LEFT

SECTIONS = (
    "Подробное описание решения",
    "Маркетинговая часть решения",
    "Бизнесовая составляющая решения",
    "Техническая проработка решения",
    "Уникальность решения",
    "Планы по развитию решения",
)


# --------------------------------------------------------------------------- данные


def facts() -> dict[str, str]:
    """facts.json прогонов пилота; число правил, актов и видов - заново из config/.

    Свод норм меняется чаще, чем пересобирается документация: слайд не должен отставать
    от config/rules.yaml.
    """
    found = json.loads((GENERATED / "facts.json").read_text(encoding="utf-8"))
    return {**found, **config_counts()}


def config_counts() -> dict[str, str]:
    """Правила, акты и виды по config/ - так же, как их считает tools/docs/annexes.py."""
    # пакет проекта нужен только здесь
    from green.domain.norms import CitationStatus  # noqa: PLC0415
    from green.infrastructure.config.repositories import (  # noqa: PLC0415
        YamlRuleBookSource,
        YamlSpeciesCatalog,
    )

    config = ROOT / "config"
    book = YamlRuleBookSource(config / "acts.yaml", config / "rules.yaml").load()
    rules = book.all_rules
    species = YamlSpeciesCatalog(config / "species.yaml").all()
    by_text = sum(
        1
        for r in rules
        if r.citation.act_id != "PROJECT" and r.citation.status is CitationStatus.VERIFIED
    )
    return {
        "rules_total": str(len(rules)),
        "rules_distance": str(len(book.distance_rules)),
        "rules_by_text": str(by_text),
        "acts": str(sum(1 for act_id in book.acts if act_id != "PROJECT")),
        "species": str(len(species)),
    }


def table_rows(markdown: str) -> list[list[str]]:
    """Строки markdown-таблицы без шапки и разделителя."""
    rows = [
        [c.strip() for c in line.strip().strip("|").split("|")]
        for line in markdown.splitlines()
        if line.strip().startswith("|")
    ]
    return [r for r in rows[2:] if r]


def num(value: str) -> float:
    return float(value.replace("\xa0", "").replace(" ", "").replace(",", "."))


SHORT_NAMES = {
    "Олимпийская деревня": "Олимпийская деревня",
    "Песчаный переулок": "Песчаный пер.",
    "3-я Парковая улица": "3-я Парковая",
    "Харьковская улица": "Харьковская",
    "улица Багрицкого": "Багрицкого",
    "Камчатская улица": "Камчатская",
    "Улица Нижние Поля": "Нижние Поля",
    "Лодочная улица": "Лодочная",
    "Измайловская площадь": "Измайловская пл.",
    "Старый Гай ул.": "Старый Гай",
    "Фрунзенская набережная": "Фрунзенская наб.",
    "Наташинский пр-д": "Наташинский пр-д",
    "Харьковский проезд": "Харьковский пр-д",
    "Куликовская улица": "Куликовская",
    "Академика Понтрягина ул.": "Ак. Понтрягина",
    "улица Берзарина": "Берзарина",
    "Малая Грузинская ул.": "Малая Грузинская",
    "Кустанайская улица": "Кустанайская",
    "2-я Прядильная ул.": "2-я Прядильная",
    "Улица Сергея Макеева": "Сергея Макеева",
}


def street_minutes(f: dict[str, str]) -> list[tuple[str, float]]:
    """Минуты прогона по улицам: таблица техники из facts.json (приложение B)."""
    return [(SHORT_NAMES[r[0]], num(r[1])) for r in table_rows(f["streets_tech_table"])]


def street_curbs() -> list[tuple[str, float, float]]:
    """Борта под зеленью, % было и стало: первая таблица приложения B."""
    text = (GENERATED / "B-streets.md").read_text(encoding="utf-8")
    first = text.split("| № | Улица |", 1)[1].split("\n\n", 1)[0]
    out = []
    for r in table_rows("| № | Улица |" + first):
        before, after = (num(v) for v in r[5].split("→"))
        out.append((SHORT_NAMES[r[1]], before, after))
    return out


# нормоконтроль принятого плана улицы Берзарина (docs/notes/24-audit.md): 785 нарушений
AUDIT = (
    ("силовой кабель", 159),
    ("бортовой камень", 104),
    ("опора освещения", 77),
    ("теплосеть", 74),
    ("водопровод", 71),
    ("кабель связи", 62),
    ("стена здания", 59),
    ("граница покрытия", 49),
    ("прочие нормы", 130),
)
AUDIT_BAD, AUDIT_ALL = 391, 554
# тесты на 20.10.2026: uv run pytest --collect-only, npx vitest run в frontend/
TESTS_BACKEND, TESTS_FRONTEND = "2 623", "1 082"


# --------------------------------------------------------------------------- примитивы


@dataclass(frozen=True)
class Style:
    font: str = MEDIUM
    size: float = 13
    color: str = BODY
    bold: bool = False


type Chunk = tuple[str, Style]
type Para = Sequence[Chunk]


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def style_run(run: object, st: Style) -> None:
    run.font.name = st.font
    run.font.size = Pt(st.size)
    run.font.bold = st.bold
    run.font.color.rgb = rgb(st.color)
    rpr = run._r.get_or_add_rPr()
    rpr.set("lang", "ru-RU")
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = rpr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rpr, qn(tag))
        el.set("typeface", st.font)


def fill_frame(
    tf: object,
    paras: Sequence[Para],
    *,
    line: float = 1.15,
    after: float = 0,
    align: PP_ALIGN = PP_ALIGN.LEFT,
    anchor: MSO_ANCHOR = MSO_ANCHOR.TOP,
    bullet: str | None = None,
) -> None:
    tf.word_wrap = True
    tf.auto_size = None
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line
        p.space_after = Pt(after)
        if bullet:
            ppr = p._p.get_or_add_pPr()
            ppr.set("marL", str(Inches(0.22)))
            ppr.set("indent", str(-Inches(0.22)))
            clr = etree.SubElement(ppr, qn("a:buClr"))
            etree.SubElement(clr, qn("a:srgbClr"), val=bullet)
            etree.SubElement(ppr, qn("a:buFont"), typeface="Arial")
            etree.SubElement(ppr, qn("a:buChar"), char="•")
        for chunk, st in para:
            if chunk == "\n":
                p.add_line_break()
                continue
            r = p.add_run()
            r.text = chunk
            style_run(r, st)


def text(
    slide: Slide,
    box: tuple[float, float, float, float],
    paras: Sequence[Para],
    **flow: object,
) -> Shape:
    x, y, w, h = box
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    fill_frame(shape.text_frame, paras, **flow)
    return shape


def para(value: str, st: Style) -> Para:
    return [(value, st)]


def alpha(shape: Shape, fill_pct: int | None = None, line_pct: int | None = None) -> None:
    """Прозрачность заливки и обводки в процентах непрозрачности."""
    sppr = shape._element.spPr
    if fill_pct is not None:
        clr = sppr.find(qn("a:solidFill")).find(qn("a:srgbClr"))
        etree.SubElement(clr, qn("a:alpha"), val=str(fill_pct * 1000))
    if line_pct is not None:
        clr = sppr.find(qn("a:ln")).find(qn("a:solidFill")).find(qn("a:srgbClr"))
        etree.SubElement(clr, qn("a:alpha"), val=str(line_pct * 1000))


def rect(
    slide: Slide,
    box: tuple[float, float, float, float],
    fill: str | None,
    *,
    line: tuple[str, float] | None = None,
    radius: float | None = None,
    shape: MSO_SHAPE = MSO_SHAPE.RECTANGLE,
) -> Shape:
    """Прямоугольник; radius - скругление в дюймах."""
    x, y, w, h = box
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else shape
    s = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    s.shadow.inherit = False
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = rgb(fill)
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = rgb(line[0])
        s.line.width = Pt(line[1])
    if radius:
        s.adjustments[0] = min(0.5, radius / min(w, h))
    return s


def circle_number(slide: Slide, x: float, y: float, d: float, label: str, *, size: float) -> None:
    c = rect(slide, (x, y, d, d), PINK, shape=MSO_SHAPE.OVAL)
    fill_frame(
        c.text_frame,
        [para(label, Style(X_BOLD, size, WHITE))],
        align=PP_ALIGN.CENTER,
        anchor=MSO_ANCHOR.MIDDLE,
    )


def arrow(slide: Slide, start: tuple[float, float], end: tuple[float, float], color: str) -> None:
    line = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(start[0]), Inches(start[1]), Inches(end[0]), Inches(end[1])
    )
    line.line.color.rgb = rgb(color)
    line.line.width = Pt(2.5)
    ln = line.line._get_or_add_ln()
    etree.SubElement(ln, qn("a:tailEnd"), type="triangle", w="med", len="med")


# --------------------------------------------------------------------------- картинки


def asset(name: str) -> Path:
    """Снимок сервиса в JPEG целиком: обрезка - свойствами кадра, файл меньше PNG вчетверо."""
    ASSETS.mkdir(parents=True, exist_ok=True)
    src = SHOTS / name
    out = ASSETS / f"{src.stem}.jpg"
    if not out.exists() or out.stat().st_mtime < src.stat().st_mtime:
        Image.open(src).convert("RGB").save(out, "JPEG", quality=86, optimize=True)
    return out


@dataclass(frozen=True)
class Crop:
    """Кадр как object-fit: cover с object-position и увеличением scale вокруг origin."""

    pos: tuple[float, float] = (0.5, 0.5)
    scale: float = 1.0
    origin: tuple[float, float] = (0.5, 0.5)


# снимки 3D-вида и фото ИИ (shots.mjs): кадр как есть
SCENE_CROP = Crop(scale=1.12)  # без шапки сайта и кнопки «Панели»
# фото ИИ: JPEG после апскейла до 4K или PNG без апскейлера - что скачал shots.mjs
PHOTO = next((n for n in ("photo-ai.jpg", "photo-ai.png") if (SHOTS / n).exists()), "photo-ai.jpg")
PHOTO_CROP = Crop()
REVIEW_CROP = Crop((0.5, 0.0))  # страница по центру окна, поля по краям срезаны
WORK_CROP = Crop((0.5, 0.5))


def picture(
    slide: Slide,
    name: str,
    box: tuple[float, float, float, float],
    crop: Crop = Crop(),  # noqa: B008 - неизменяемый dataclass
    *,
    dark: bool = False,
    radius: float = 0.125,
    shadow: bool = True,
    border: str | None = None,
) -> None:
    path = asset(name)
    x, y, w, h = box
    iw, ih = Image.open(path).size
    frame = w / h
    if iw / ih > frame:
        vw, vh = ih * frame, ih
        left, top = (iw - vw) * crop.pos[0], 0.0
    else:
        vw, vh = iw, iw / frame
        left, top = 0.0, (ih - vh) * crop.pos[1]
    s = crop.scale
    x0 = left + crop.origin[0] * (1 - 1 / s) * vw
    y0 = top + crop.origin[1] * (1 - 1 / s) * vh
    pic = slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w), Inches(h))
    pic.crop_left = x0 / iw
    pic.crop_right = 1 - (x0 + vw / s) / iw
    pic.crop_top = y0 / ih
    pic.crop_bottom = 1 - (y0 + vh / s) / ih
    sppr = pic._element.spPr
    geom = sppr.find(qn("a:prstGeom"))
    geom.set("prst", "roundRect")
    av = geom.find(qn("a:avLst"))
    if av is None:
        av = etree.SubElement(geom, qn("a:avLst"))
    etree.SubElement(av, qn("a:gd"), name="adj", fmla=f"val {round(radius / min(w, h) * 100000)}")
    if border:
        pic.line.color.rgb = rgb(border)
        pic.line.width = Pt(2.25)
    if shadow:
        sppr.append(
            parse_xml(
                f'<a:effectLst {nsdecls("a")}><a:outerShdw blurRad="330000" dist="120000" '
                f'dir="5400000" algn="t" rotWithShape="0"><a:srgbClr val="{"000000" if dark else NIGHT}">'
                f'<a:alpha val="{40000 if dark else 26000}"/></a:srgbClr></a:outerShdw></a:effectLst>'
            )
        )


def tag(slide: Slide, x: float, y: float, label: str) -> None:
    """Подпись поверх снимка: тёмная плашка с белым текстом."""
    w = 0.35 + len(label) * 0.088
    t = rect(slide, (x, y, w, 0.3), INK, radius=0.15)
    alpha(t, 82)
    fill_frame(
        t.text_frame,
        [para(label, Style(SEMI, 10, WHITE))],
        align=PP_ALIGN.CENTER,
        anchor=MSO_ANCHOR.MIDDLE,
    )
    t.text_frame.word_wrap = False


# --------------------------------------------------------------------------- каркас слайда


def set_background(slide: Slide, image: Path) -> None:
    """Фон слайда картинкой шаблона: не выделяется и не сдвигается при правке."""
    _, rid = slide.part.get_or_add_image_part(str(image))
    bg = parse_xml(
        f"<p:bg {nsdecls('p', 'a', 'r')}><p:bgPr>"
        f'<a:blipFill dpi="0" rotWithShape="1"><a:blip r:embed="{rid}"/><a:srcRect/>'
        "<a:stretch><a:fillRect/></a:stretch></a:blipFill><a:effectLst/></p:bgPr></p:bg>"
    )
    slide._element.cSld.insert(0, bg)


def slide_number(slide: Slide, *, dark: bool) -> None:
    shape = slide.shapes.add_textbox(Inches(12.3), Inches(7.0), Inches(0.5), Inches(0.28))
    tf = shape.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    fld = etree.SubElement(
        p._p, qn("a:fld"), id="{B6F15528-21DE-4FAA-801E-634DDDAF4B2B}", type="slidenum"
    )
    rpr = etree.SubElement(fld, qn("a:rPr"), lang="ru-RU", sz="1000", b="1")
    fill = etree.SubElement(rpr, qn("a:solidFill"))
    etree.SubElement(fill, qn("a:srgbClr"), val=WHITE if dark else PLUM)
    etree.SubElement(rpr, qn("a:latin"), typeface=SEMI)
    etree.SubElement(fld, qn("a:t")).text = "#"


def new_slide(prs: Deck, section: int, heading: str, *, bg: str = "light") -> Slide:
    """Слайд раздела section (1-6): фон шаблона, метка раздела, заголовок, номер."""
    layout = next(lay for lay in prs.slide_layouts if lay.name == "Пустой")
    slide = prs.slides.add_slide(layout)
    set_background(slide, BG / f"{bg}.png")
    dark = bg != "light"
    label = f"{section:02d}   {SECTIONS[section - 1].upper()}"
    chip = rect(slide, (LEFT, 0.42, 0.45 + len(label) * 0.098, 0.36), PINK, radius=0.18)
    fill_frame(
        chip.text_frame,
        [para(label, Style(SEMI, 10.5, WHITE))],
        align=PP_ALIGN.CENTER,
        anchor=MSO_ANCHOR.MIDDLE,
    )
    chip.text_frame.word_wrap = False
    text(
        slide,
        (LEFT, 0.98, WIDTH, 1.0),
        [para(heading, Style(X_BOLD, 30, WHITE if dark else NIGHT))],
        line=1.0,
    )
    slide_number(slide, dark=dark)
    return slide


def notes(slide: Slide, say: str) -> None:
    slide.notes_slide.notes_text_frame.text = say


def foot(slide: Slide, value: str, *, dark: bool, y: float = 6.55, w: float = WIDTH) -> None:
    text(slide, (LEFT, y, w, 0.6), [para(value, Style(MEDIUM, 12, MIST if dark else BODY))])


# --------------------------------------------------------------------------- графики


def style_chart(chart: Chart, *, dark: bool, size: float = 11) -> None:
    chart.font.name = MEDIUM
    chart.font.size = Pt(size)
    chart.font.color.rgb = rgb(MIST if dark else BODY)
    chart.has_title = False


def paint(series: object, color: str) -> None:
    series.format.fill.solid()
    series.format.fill.fore_color.rgb = rgb(color)
    series.format.line.fill.background()


def hbar_chart(
    slide: Slide,
    box: tuple[float, float, float, float],
    categories: Sequence[str],
    series: Sequence[tuple[str, Sequence[float], str]],
    *,
    dark: bool = False,
    number_format: str = "0",
    size: float = 11,
    legend: bool = True,
) -> Chart:
    data = CategoryChartData(number_format=number_format)
    data.categories = list(categories)
    for name, values, _ in series:
        data.add_series(name, list(values))
    x, y, w, h = box
    chart = slide.shapes.add_chart(
        XL_CHART_TYPE.BAR_CLUSTERED, Inches(x), Inches(y), Inches(w), Inches(h), data
    ).chart
    style_chart(chart, dark=dark, size=size)
    plot = chart.plots[0]
    plot.gap_width = 40
    plot.overlap = 0
    plot.vary_by_categories = False
    plot.has_data_labels = True
    labels = plot.data_labels
    labels.number_format = number_format
    labels.number_format_is_linked = False
    labels.position = XL_LABEL_POSITION.OUTSIDE_END
    labels.font.size = Pt(size)
    labels.font.bold = True
    labels.font.color.rgb = rgb(WHITE if dark else NIGHT)
    for (_, _, color), s in zip(series, plot.series, strict=True):
        paint(s, color)
    cat = chart.category_axis
    cat.reverse_order = True
    cat.has_major_gridlines = False
    cat.format.line.fill.background()
    cat.tick_labels.font.size = Pt(size)
    cat.tick_labels.font.color.rgb = rgb(MIST if dark else BODY)
    val = chart.value_axis
    val.visible = False
    val.has_major_gridlines = False
    val.minimum_scale = 0
    chart.has_legend = len(series) > 1 and legend
    if chart.has_legend:
        chart.legend.position = XL_LEGEND_POSITION.TOP
        chart.legend.include_in_layout = False
        chart.legend.font.size = Pt(size)
        chart.legend.font.color.rgb = rgb(WHITE if dark else BODY)
    return chart


# --------------------------------------------------------------------------- 01 описание


def s_problem(prs: Deck, f: dict[str, str]) -> None:
    s = new_slide(prs, 1, "Под газоном идут кабели и трубы")
    picture(s, "utilities.png", (LEFT, 2.1, 6.95, 4.45), Crop((0.5, 0.5), 1.35, (0.5, 0.45)))
    x, w = 8.0, 4.67
    text(
        s,
        (x, 2.1, w, 2.0),
        [
            para(
                "Дерево нельзя посадить ближе 2 м к силовому кабелю, 1,5 м к газопроводу, "
                "5 м к стене дома (СП 42.13330.2016, табл. 9.1).",
                Style(SEMI, 15, INK),
            )
        ],
        line=1.2,
    )
    text(
        s,
        (x, 3.55, w, 1.1),
        [
            para(
                "Проектировщик сверяет каждое место с сетями на чертеже вручную. Заказчик "
                "назвал разбор сетей самым трудоёмким этапом проекта.",
                Style(MEDIUM, 12.5, BODY),
            )
        ],
        line=1.25,
    )
    for i, (value, label) in enumerate(
        (("14 316", "мест-кандидатов вдоль бортов Кустанайской"), ("1 639", "проходят все нормы"))
    ):
        y = 4.85 + i * 0.85
        text(s, (x, y, 2.0, 0.6), [para(value, Style(X_BOLD, 26, PINK))])
        text(s, (x + 2.05, y + 0.1, w - 2.05, 0.6), [para(label, Style(MEDIUM, 12, BODY))])
    text(
        s,
        (LEFT, 6.7, 6.95, 0.3),
        [para("Подземные сети Кустанайской улицы, слои плана скрыты", Style(MEDIUM, 10, MUTE))],
    )
    notes(
        s,
        "Задача: посадить деревья и кусты по нормам отступов от сетей, зданий и бортов. "
        f"На Кустанайской (прогон 20.10.2026) из 14 316 мест вдоль бортов нормы проходят 1 639. "
        f"Всего в своде {f['rules_distance']} норм расстояний.",
    )


def s_solution(prs: Deck, f: dict[str, str]) -> None:
    s = new_slide(prs, 1, "DXF на входе, план и объяснения на выходе", bg="dark")
    w, h, gap = 3.5, 2.7, 0.75
    frames = (
        (
            "flow-before.png",
            Crop((0.3, 0.5), 2.0, (0.25, 0.3)),
            "DXF улицы",
            "подоснова Мосгеотреста и сети; по желанию слои ГИС, кадастр, перечётная ведомость",
        ),
        (
            "flow-after.png",
            Crop((0.3, 0.5), 2.0, (0.25, 0.3)),
            "План посадок",
            "слои GREEN_*: деревья, кусты, газоны, цветники, отказы, зоны допустимости",
        ),
        (
            "explain-card.png",
            Crop((0.5, 0.0)),
            "Объяснение каждой посадки",
            "замер, норма, акт и пункт",
        ),
    )
    for i, (name, crop, head, sub) in enumerate(frames):
        x = LEFT + i * (w + gap)
        picture(s, name, (x, 2.1, w, h), crop, dark=True)
        text(
            s,
            (x, 2.1 + h + 0.2, w, 1.1),
            [para(head, Style(SEMI, 15, WHITE)), para(sub, Style(MEDIUM, 11.5, MIST))],
            line=1.2,
            after=3,
        )
        if i < len(frames) - 1:
            arrow(s, (x + w + 0.15, 2.1 + h / 2), (x + w + gap - 0.15, 2.1 + h / 2), ROSE)
    foot(
        s,
        f"Исходные слои чертежа не меняются: на {f['streets_run']} улицах сверено "
        f"{f['entities']} объектов, изменено {f['entities_changed']}. У каждой посадки в DXF "
        "атрибуты вида и нормы.",
        dark=True,
    )
    notes(
        s, "Вход - DXF улицы, выход - тот же DXF со слоями GREEN_* и объяснения по каждой посадке."
    )


def s_steps(prs: Deck, f: dict[str, str]) -> None:
    s = new_slide(prs, 1, "Семь шагов от DXF до DXF")
    steps = (
        (
            "Чтение без потерь",
            "DXF, DWG через ODA; каждый примитив учтён, пунктир съёмки собран в линии",
        ),
        (
            "Классы слоёв и знаков",
            "сети, борта, здания, деревья; незнакомый слой - по словам имени и геометрии",
        ),
        ("Карта покрытий", "газон, тротуар и проезд: где под ямой грунт, в том числе по подписям"),
        (
            "Места посадки",
            "аллея вдоль борта, группы деревьев с газоном между ними, ряды кустарника, цветники",
        ),
        (
            "Вид под место",
            f"{f['species']} видов, MILP по условиям места и квотам; светолюбивые - не в тени домов",
        ),
        (
            "Независимая проверка",
            f"{f['rules_total']} правил заново по готовому плану, нарушений {f['violations']}",
        ),
        (
            "Запись и объяснения",
            "слои GREEN_*, ведомость, отчёт с пунктами актов, записка DOCX",
        ),
    )
    gap = 0.16
    w = (WIDTH - gap * 6) / 7
    for i, (head, sub) in enumerate(steps):
        x = LEFT + i * (w + gap)
        circle_number(s, x, 2.3, 0.56, str(i + 1), size=16)
        if i < len(steps) - 1:
            dash = s.shapes.add_connector(
                MSO_CONNECTOR.STRAIGHT,
                Inches(x + 0.66),
                Inches(2.58),
                Inches(x + w + gap - 0.1),
                Inches(2.58),
            )
            dash.line.color.rgb = rgb(BLUSH)
            dash.line.width = Pt(2)
        text(s, (x, 3.1, w, 0.75), [para(head, Style(X_BOLD, 13.5, NIGHT))], line=1.05)
        text(s, (x, 3.9, w, 2.2), [para(sub, Style(MEDIUM, 11.5, BODY))], line=1.25)
    card = rect(s, (LEFT, 5.75, WIDTH, 0.95), WHITE, line=(BLUSH, 1.25), radius=0.16)
    card.text_frame.margin_left = card.text_frame.margin_right = Inches(0.3)
    fill_frame(
        card.text_frame,
        [
            [
                ("Алгоритм, а не обучение. ", Style(SEMI, 12.5, PINK)),
                (
                    "Норма - порог расстояния, а принятые планы нарушают её в 391 посадке из 554: "
                    "учиться на них нельзя. Места и виды выбирает задача MILP (HiGHS). Локальная "
                    "языковая модель только подсказывает класс незнакомого слоя, план она не меняет.",
                    Style(MEDIUM, 12.5, BODY),
                ),
            ]
        ],
        anchor=MSO_ANCHOR.MIDDLE,
        line=1.2,
    )
    card.text_frame.margin_left = card.text_frame.margin_right = Inches(0.3)
    notes(
        s, "Семь шагов конвейера. Ключевое: генерация и проверка разделены, проверка независимая."
    )


def s_explain(prs: Deck, f: dict[str, str]) -> None:
    s = new_slide(prs, 1, "У каждой посадки есть пункт нормы")
    picture(s, "explain-map.png", (LEFT, 1.95, 4.85, 4.7), Crop((0.3, 0.5), 1.3, (0.5, 0.5)))
    picture(s, "explain-card.png", (5.75, 1.95, 3.25, 4.7), Crop((0.5, 0.0)))
    x, w = 9.35, 3.32
    text(
        s,
        (x, 2.0, w, 1.6),
        [
            para(
                "Выноска на карте показывает замер: «теплосеть 3,47 ≥ 2,50». В карточке - "
                "правило, акт, пункт и дословная цитата.",
                Style(SEMI, 14, INK),
            )
        ],
        line=1.2,
    )
    text(
        s,
        (x, 3.75, w, 3.0),
        [
            para(
                "Отказ объясняется так же: «до силового кабеля 0,51 м < 2,00 м». "
                "Правило R-POWER-TREE-001, основание - СП 42.13330.2016, п. 9.6, табл. 9.1.",
                Style(MEDIUM, 12, BODY),
            ),
            para(
                f"Свод: {f['rules_total']} правил, {f['acts']} актов. У {f['rules_by_text']} "
                f"основание сверено по тексту акта. Текст объяснения собирается из трассы "
                "проверок без языковой модели.",
                Style(MEDIUM, 12, BODY),
            ),
        ],
        line=1.25,
        after=10,
    )
    notes(s, "Интерпретируемость: у каждой посадки и отказа замер, правило, акт, пункт, цитата.")


# --------------------------------------------------------------------------- 02 маркетинг


def s_users(prs: Deck, f: dict[str, str]) -> None:
    s = new_slide(prs, 2, "Кому нужен сервис")
    roles = (
        (
            "Проектировщик благоустройства",
            "Подаёт DXF подосновы и сетей, получает черновик плана с проверенными отступами, "
            "правит его в браузере и доводит в своём CAD.",
            f"медиана {f['minutes_median']} мин на улицу",
        ),
        (
            "Специалист ДПиООС",
            "Проверяет отступы и обоснования: у каждой посадки акт, пункт, расстояние и запас. "
            "Теми же правилами проверяет план подрядчика.",
            "нормоконтроль green audit",
        ),
        (
            "Участник согласования",
            "Сверяет объёмы озеленения и баланс «было - стало»: ведомость и пояснительная "
            "записка собраны из того же плана, что и DXF.",
            "ведомость Excel и записка DOCX",
        ),
    )
    gap = 0.3
    w = (WIDTH - 2 * gap) / 3
    for i, (head, body, metric) in enumerate(roles):
        x = LEFT + i * (w + gap)
        rect(s, (x, 1.95, w, 2.85), WHITE, line=(BLUSH, 1.25), radius=0.18)
        circle_number(s, x + 0.3, 2.22, 0.5, str(i + 1), size=14)
        text(
            s,
            (x + 0.95, 2.22, w - 1.2, 0.55),
            [para(head, Style(X_BOLD, 14, NIGHT))],
            anchor=MSO_ANCHOR.MIDDLE,
            line=1.0,
        )
        text(s, (x + 0.3, 2.95, w - 0.6, 1.3), [para(body, Style(MEDIUM, 11.5, BODY))], line=1.25)
        text(s, (x + 0.3, 4.33, w - 0.6, 0.35), [para(metric, Style(SEMI, 12, PINK))])
    text(
        s,
        (LEFT, 5.05, WIDTH, 0.35),
        [para("Что сказал заказчик на сессии вопросов", Style(SEMI, 13, PLUM))],
    )
    quotes = (
        ("№ 1", "причина возврата проектов - нарушение отступов, № 2 - неверные объёмы озеленения"),
        ("0", "попыток автоматизировать проектирование озеленения у заказчика не было"),
        ("DXF", "проектировщики работают в AutoCAD без плагинов: результат - обычные слои чертежа"),
        ("до 60 мин", "ориентир на генерацию плана улицы, до 30 мин на подбор мест"),
    )
    qw = (WIDTH - 3 * 0.25) / 4
    for i, (big, label) in enumerate(quotes):
        x = LEFT + i * (qw + 0.25)
        text(s, (x, 5.45, qw, 0.45), [para(big, Style(X_BOLD, 20, PINK))])
        text(s, (x, 5.95, qw, 0.9), [para(label, Style(MEDIUM, 10.5, BODY))], line=1.2)
    notes(
        s,
        "Три пользователя из ТЗ, у каждого свой вход. Ниже - ответы заказчика (docs/notes/15): "
        "главная боль - отступы, автоматизации не было, нужен обычный DXF.",
    )


def s_audit(prs: Deck, f: dict[str, str]) -> None:
    s = new_slide(prs, 2, "Сервис проверяет и чужие планы")
    data = CategoryChartData(number_format="0")
    data.categories = ["с нарушениями", "без нарушений"]
    data.add_series("посадки", [AUDIT_BAD, AUDIT_ALL - AUDIT_BAD])
    chart = s.shapes.add_chart(
        XL_CHART_TYPE.DOUGHNUT, Inches(LEFT), Inches(1.95), Inches(3.3), Inches(3.3), data
    ).chart
    style_chart(chart, dark=False)
    chart.has_legend = False
    plot = chart.plots[0]
    plot.vary_by_categories = True
    for point, color in zip(plot.series[0].points, (PINK, BLUSH), strict=True):
        point.format.fill.solid()
        point.format.fill.fore_color.rgb = rgb(color)
        point.format.line.fill.background()
    hole = plot._element.find(qn("c:holeSize"))
    if hole is None:
        hole = etree.SubElement(plot._element, qn("c:holeSize"))
    hole.set("val", "68")
    share = round(AUDIT_BAD / AUDIT_ALL * 100)
    text(
        s,
        (LEFT, 2.98, 3.3, 1.2),
        [
            para(f"{share}%", Style(X_BOLD, 34, PINK)),
            para("посадок", Style(SEMI, 10, BODY)),
            para("с нарушениями", Style(SEMI, 10, BODY)),
        ],
        align=PP_ALIGN.CENTER,
        line=1.0,
    )
    text(
        s,
        (LEFT, 5.45, 4.2, 1.4),
        [
            para(
                f"В принятом плане улицы Берзарина {AUDIT_BAD} из {AUDIT_ALL} посадок нарушают "
                "отступы; нормоконтроль занял 162 с.",
                Style(MEDIUM, 12, BODY),
            ),
            para(
                f"В плане сервиса для той же улицы нарушений {f['violations']}.",
                Style(SEMI, 12, NIGHT),
            ),
        ],
        line=1.2,
        after=6,
    )
    x = 5.35
    text(
        s,
        (x, 1.95, RIGHT - x, 0.4),
        [
            para(
                "785 нарушений по видам объектов: у посадки их бывает несколько",
                Style(SEMI, 13, NIGHT),
            )
        ],
    )
    chart = hbar_chart(
        s,
        (x, 2.35, RIGHT - x, 4.4),
        [a for a, _ in AUDIT],
        [("нарушения", [v for _, v in AUDIT], PINK)],
    )
    last = chart.plots[0].series[0].points[len(AUDIT) - 1]
    last.format.fill.solid()
    last.format.fill.fore_color.rgb = rgb(LAV)
    notes(
        s,
        "Второй продукт - нормоконтроль: green audit проверяет план проектировщика теми же "
        "правилами и пишет, сколько метров не хватает.",
    )


def s_visual(prs: Deck, f: dict[str, str]) -> None:  # noqa: ARG001 - общая подпись слайдов
    s = new_slide(prs, 2, "Улицу видно до начала работ", bg="dark")
    # 3D-вид в ТЗ не просили, фото ИИ - бонус ТЗ: жюри видит это до того, как читает слайд.
    # Плашка - на уровне заголовка у правого края: верхнюю полосу занимают логотипы шаблона.
    extra = "СВЕРХ ТЗ: 3D-ВИД  ·  БОНУС ТЗ: ФОТО ИИ"
    w = 0.45 + len(extra) * 0.098
    chip = rect(s, (RIGHT - w, 1.2, w, 0.4), LAV, radius=0.2)
    fill_frame(
        chip.text_frame,
        [para(extra, Style(SEMI, 10.5, WHITE))],
        align=PP_ALIGN.CENTER,
        anchor=MSO_ANCHOR.MIDDLE,
    )
    chip.text_frame.word_wrap = False
    picture(s, "scene3d.png", (LEFT, 1.95, 7.3, 4.1), SCENE_CROP, dark=True)
    tag(s, LEFT + 0.2, 1.95 + 4.1 - 0.5, "3D-вид по координатам плана")
    x, w = 8.25, RIGHT - 8.25
    picture(s, PHOTO, (x, 1.95, w, 1.95), PHOTO_CROP, dark=True)
    tag(s, x + 0.15, 1.95 + 1.95 - 0.45, "фото ИИ по этому кадру")
    picture(s, "scene-night.png", (x, 4.1, w, 1.95), SCENE_CROP, dark=True)
    tag(s, x + 0.15, 4.1 + 1.95 - 0.45, "тот же кадр в 20:30, лето")
    foot(
        s,
        "Покрытия, борта, ограды и площадки - по линиям чертежа, этажность домов - по подписям, "
        "кроны - по виду и возрасту, солнце - по часу над Москвой. Фото по кадру рисует "
        "Qwen-Image-2.1-Turbo на своей видеокарте, без внешних сервисов.",
        dark=True,
        y=6.25,
    )
    notes(
        s,
        "3D-вид сделан сверх ТЗ, фото ИИ - бонус ТЗ; оба из того же плана. Посадки стоят там же, где в DXF; по тротуарам "
        "ходят люди, на площадках дети. Время суток и сезон меняются на пульте, ночью горят "
        "фонари. Фото по кадру: режим «по плану» не дорисовывает деревья, которых нет в плане. "
        "Время: на RTX 5070 кадр 1536 x 864 - 85 с, через API с апскейлом до 4K - 94 с "
        "(docs/notes/40); фото на слайде - 1536 x 864 без апскейла.",
    )


# --------------------------------------------------------------------------- 03 бизнес


def s_city(prs: Deck, f: dict[str, str]) -> None:
    s = new_slide(prs, 3, "Эффект внедрения для города")
    cards = (
        (
            f"{f['minutes_median']} мин",
            "на генерацию плана улицы",
            f"Медиана по {f['streets_done']} улицам пилота, от {f['minutes_min']} "
            f"до {f['minutes_max']} мин.",
            "Ориентир заказчика - до 60 мин.",
        ),
        (
            f["rules_total"],
            "оцифрованных правил",
            f"{f['acts']} актов, в том числе 743-ПП и СП 42.13330; у {f['rules_by_text']} "
            "основание сверено по тексту.",
            f"{f['violations']} нарушений в планах {f['streets_done']} улиц.",
        ),
        (
            "2,7 мин",
            "на нормоконтроль плана улицы",
            "Входной контроль проектов подрядчиков: пункт акта и недобор в метрах.",
            "Нарушение отступов - главная причина возврата проектов.",
        ),
    )
    gap = 0.3
    w = (WIDTH - 2 * gap) / 3
    y, h = 1.95, 2.6
    for i, (big, sub, first, second) in enumerate(cards):
        x = LEFT + i * (w + gap)
        rect(s, (x, y, w, h), WHITE, line=(BLUSH, 1.25), radius=0.18)
        inner = w - 0.6
        text(s, (x + 0.3, y + 0.22, inner, 0.65), [para(big, Style(X_BOLD, 32, PINK))])
        text(s, (x + 0.3, y + 0.88, inner, 0.35), [para(sub, Style(SEMI, 13, NIGHT))])
        text(
            s,
            (x + 0.3, y + 1.3, inner, 1.1),
            [para(first, Style(MEDIUM, 11, BODY)), para(second, Style(MEDIUM, 11, MUTE))],
            line=1.2,
            after=5,
        )
    py, ph = 4.8, 2.0
    rect(s, (LEFT, py, WIDTH, ph), NIGHT, radius=0.18)
    text(
        s,
        (LEFT + 0.35, py + 0.22, 6.0, 0.35),
        [para("Системные преимущества", Style(SEMI, 13, ROSE))],
    )
    points = (
        (
            "Гибкость и актуальность",
            "Нормативы вынесены в конфигурацию: смена редакции или региональных правил без "
            "переписывания кода. СП 42.13330 2016 и 2026 уже в своде.",
        ),
        (
            "Безопасность данных (on-premise)",
            "Работа в закрытом контуре заказчика: геоданные не уходят на внешние серверы. "
            "Расчёту не нужны сеть и GPU, фото и подсказки модели - на своей видеокарте.",
        ),
        (
            "Проверяемость решений",
            "Каждая посадка и каждый отказ со ссылкой на акт, пункт и дословную цитату: "
            "согласующий проверяет решение, а не верит ему.",
        ),
    )
    cgap = 0.35
    cw = (WIDTH - 0.7 - 2 * cgap) / 3
    for i, (head, body) in enumerate(points):
        x = LEFT + 0.35 + i * (cw + cgap)
        text(s, (x, py + 0.65, cw, 0.35), [para(head, Style(SEMI, 13, WHITE))])
        text(s, (x, py + 1.02, cw, 0.95), [para(body, Style(MEDIUM, 10.5, MIST))], line=1.2)
    notes(
        s,
        f"Эффект для города: улица за минуты вместо ручного разбора сетей, {f['rules_total']} "
        "правил вместо ручной сверки, нормоконтроль подрядчика за 2,7 минуты. Внизу - три свойства, "
        "которые важны заказчику: нормы меняются без кода, данные не уходят из контура, "
        "каждое решение проверяемо по пункту акта.",
    )


def s_effect(prs: Deck, f: dict[str, str]) -> None:
    s = new_slide(prs, 3, f"Тень и защита от пыли на {f['effect_streets']} улицах", bg="dark")
    top = sorted(street_curbs(), key=lambda r: r[2] - r[1], reverse=True)[:7]
    text(
        s,
        (LEFT, 1.9, 7.0, 0.6),
        [
            [
                ("Борта под кронами и кустарником, %: ", Style(SEMI, 13, WHITE)),
                ("было", Style(SEMI, 13, LAV)),
                (" и ", Style(SEMI, 13, WHITE)),
                ("с планом сервиса", Style(SEMI, 13, ROSE)),
            ],
            para("7 улиц с наибольшим приростом", Style(MEDIUM, 11, MIST)),
        ],
    )
    hbar_chart(
        s,
        (LEFT, 2.55, 7.0, 4.25),
        [r[0] for r in top],
        [("было", [r[1] for r in top], LAV), ("с планом сервиса", [r[2] for r in top], PINK)],
        dark=True,
        number_format="0.0",
        size=10.5,
        legend=False,
    )
    stats = (
        (
            f"{f['trees']} и {f['shrubs']}",
            f"деревьев и кустов в планах {f['streets_done']} улиц пилота",
        ),
        (
            f["effect_curb"].replace(" длины бортов", ""),
            f"длины бортов под кронами и кустарником, прирост {f['effect_curb_m']}",
        ),
        (f["effect_canopy_m2"], f"взрослых крон: {f['effect_canopy']}"),
        (f["effect_tiers"], "деревьев с кустарником под кроной"),
    )
    x, w = 8.2, RIGHT - 8.2
    for i, (big, label) in enumerate(stats):
        y = 1.95 + i * 1.18
        text(s, (x, y, w, 0.5), [para(big, Style(X_BOLD, 22, ROSE))])
        text(s, (x, y + 0.48, w, 0.6), [para(label, Style(MEDIUM, 11, MIST))], line=1.15)
    notes(
        s,
        "Эффект для города: тень, пылезащита, ярусность - то, что заказчик назвал признаками хорошего плана.",
    )


# --------------------------------------------------------------------------- 04 техника


def node(
    slide: Slide, box: tuple[float, float, float, float], head: str, sub: str, *, core: bool = False
) -> None:
    n = rect(slide, box, PINK if core else WHITE, line=None if core else (WHITE, 1.0), radius=0.14)
    if not core:
        alpha(n, 10, 30)
    tf = n.text_frame
    fill_frame(
        tf,
        [
            para(head, Style(SEMI, 14, WHITE)),
            para(sub, Style(MEDIUM, 10.5, WHITE if core else MIST)),
        ],
        anchor=MSO_ANCHOR.MIDDLE,
        line=1.15,
        after=2,
    )
    tf.margin_left = tf.margin_right = Inches(0.2)


def s_deploy(prs: Deck, f: dict[str, str]) -> None:
    s = new_slide(prs, 4, "Разворачивается одной командой", bg="dark")
    y0, h, gap = 1.95, 0.8, 0.2
    for i, (head, sub) in enumerate(
        (
            ("CLI", "green run, audit, verify"),
            ("HTTP API", "OpenAPI 3.1, Swagger без сети"),
            ("Веб-интерфейс", "React: план, правка, 3D и фото"),
        )
    ):
        node(s, (LEFT, y0 + i * (h + gap), 2.8, h), head, sub)
    mid = y0 + (3 * h + 2 * gap) / 2
    arrow(s, (LEFT + 2.95, mid), (LEFT + 3.55, mid), ROSE)
    node(
        s,
        (LEFT + 3.7, mid - 0.55, 2.95, 1.1),
        "Одно ядро",
        "CLI, API и веб вызывают один сценарий PlanSite",
        core=True,
    )
    code = rect(s, (LEFT, 4.95, 6.65, 0.85), "000000", radius=0.12)
    alpha(code, 35)
    fill_frame(
        code.text_frame,
        [
            para("docker compose up --build", Style(MONO, 12, WHITE, bold=True)),
            para(
                "green run улица.dxf --profile barriers --set spacing_m=6",
                Style(MONO, 12, WHITE, bold=True),
            ),
        ],
        anchor=MSO_ANCHOR.MIDDLE,
        line=1.3,
    )
    code.text_frame.margin_left = Inches(0.25)
    text(
        s,
        (LEFT, 5.95, 6.65, 0.9),
        [
            para(
                "Образ на Ubuntu 26.04: по ТЗ МосТех.ОС близка к Ubuntu; ODA File Converter - "
                f"аргументом сборки. {f['params']} параметров, 5 профилей; {TESTS_BACKEND} теста "
                f"бэкенда и {TESTS_FRONTEND} фронтенда. Ollama для подсказок - профиль llm в "
                "compose. Манифесты Kubernetes проверены kubeconform.",
                Style(MEDIUM, 11.5, MIST),
            )
        ],
        line=1.2,
    )
    picture(s, "swagger.png", (7.85, 1.95, RIGHT - 7.85, 4.85), Crop((0.0, 0.0)), dark=True)
    notes(s, "Архитектура: гексагональное ядро, три тонкие обёртки. Swagger без интернета.")


def s_speed(prs: Deck, f: dict[str, str]) -> None:
    s = new_slide(prs, 4, f"{f['streets_done']} улиц пилота без ручных правок")
    rows = sorted(street_minutes(f), key=lambda r: r[1])
    text(
        s,
        (LEFT, 1.9, 8.7, 0.6),
        [
            para(
                f"Минуты прогона по улицам, медиана {f['minutes_median']}", Style(SEMI, 13, NIGHT)
            ),
            para(
                "Пунктир - ориентиры заказчика: 30 мин на подбор мест, 60 мин на генерацию",
                Style(MEDIUM, 11, BODY),
            ),
        ],
    )
    data = CategoryChartData(number_format="0.0")
    data.categories = [r[0] for r in rows]
    data.add_series("минуты", [r[1] for r in rows])
    chart = s.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(LEFT), Inches(2.5), Inches(8.7), Inches(4.4), data
    ).chart
    style_chart(chart, dark=False, size=9)
    chart.has_legend = False
    plot = chart.plots[0]
    plot.gap_width = 45
    plot.has_data_labels = True
    labels = plot.data_labels
    labels.number_format = "0.0"
    labels.number_format_is_linked = False
    labels.position = XL_LABEL_POSITION.OUTSIDE_END
    labels.font.size = Pt(9)
    labels.font.bold = True
    labels.font.color.rgb = rgb(NIGHT)
    paint(plot.series[0], PINK)
    val = chart.value_axis
    val.minimum_scale = 0
    val.maximum_scale = 60
    val.major_unit = 30
    val.has_major_gridlines = True
    grid = val.major_gridlines.format.line
    grid.color.rgb = rgb(PLUM)
    grid.width = Pt(1.25)
    grid.dash_style = MSO_LINE_DASH_STYLE.DASH
    val.format.line.fill.background()
    val.tick_labels.font.size = Pt(9)
    val.tick_labels.number_format = "0"
    val.tick_labels.number_format_is_linked = False
    cat = chart.category_axis
    cat.format.line.color.rgb = rgb(BLUSH)
    cat.tick_labels.font.size = Pt(9)
    body_pr = cat._element.find(qn("c:txPr")).find(qn("a:bodyPr"))
    body_pr.set("rot", "-2700000")
    body_pr.set("vert", "horz")
    x, w = 9.75, RIGHT - 9.75
    stats = (
        (
            f"{f['streets_done']} из {len(rows)}",
            "улиц каталога с планом; Фрунзенская собрана из PDF-генплана, DWG у неё нет",
        ),
        (f["violations"], "нарушений норм в выданных планах по независимой проверке"),
        (
            f"{f['entities_changed']} из {f['entities']}",
            "объектов исходников изменено после записи",
        ),
    )
    for i, (big, label) in enumerate(stats):
        y = 2.0 + i * 1.6
        text(s, (x, y, w, 0.55), [para(big, Style(X_BOLD, 24 if len(big) < 10 else 17, PINK))])
        text(s, (x, y + 0.6, w, 0.9), [para(label, Style(MEDIUM, 11, BODY))], line=1.2)
    notes(
        s,
        "Время по всем улицам пилота, строгий профиль. Каждая улица - один прогон без ручных "
        "правок. Фрунзенская набережная в материалах пилота есть только листами PDF: "
        "tools/pdf_to_dxf.py собрал из пяти листов чертёж в метрах.",
    )


def s_inputs(prs: Deck, f: dict[str, str]) -> None:  # noqa: ARG001 - общая подпись слайдов
    s = new_slide(prs, 4, "Чужой чертёж и данные города")
    picture(s, "review.png", (LEFT, 1.95, 6.3, 4.85), REVIEW_CROP)
    text(
        s,
        (LEFT, 6.88, 6.3, 0.3),
        [
            para(
                "Уточнение объектов чертежа: класс каждого слоя и основание",
                Style(MEDIUM, 10, MUTE),
            )
        ],
    )
    x, w = 7.35, RIGHT - 7.35
    items = (
        (
            "Чертёж жюри",
            "236287.dxf проходит без ручных правок: слова слоёв паспорта ОГХ, единицы без "
            "заголовка, пунктир MicroStation собирается в линии.",
        ),
        (
            "Незнакомый слой",
            "Класс - по словам имени и по геометрии. Локальная модель в Ollama (gemma4) "
            "подсказывает в таблице уточнения, решает человек.",
        ),
        (
            "Слои ГИС",
            "GeoJSON и SHP. Охранные и природные территории задают тип территории по 369-ПП, "
            "охранные зоны газа и ЛЭП из кадастра стесняют посадку.",
        ),
        (
            "Кадастр",
            "Участок задаёт назначение места. Яма на частном участке - на согласование с "
            "правообладателем, на участке дома - с собранием собственников.",
        ),
    )
    y = 1.95
    for head, body in items:
        text(
            s,
            (x, y, w, 1.15),
            [para(head, Style(X_BOLD, 13, NIGHT)), para(body, Style(MEDIUM, 11.5, BODY))],
            line=1.2,
            after=3,
        )
        y += 1.2
    notes(
        s,
        "Устойчивость к входу: чертёж жюри 236287 считается без JSON-переопределений; "
        "Фрунзенская собрана из PDF. Подсказки LLM - совет в таблице уточнения, класс и план они "
        "не меняют. Кадастр: ГрК РФ ст. 1 п. 12, ЖК РФ ст. 44 - правила R-PARCEL-* и R-MKD-*.",
    )


def s_workspace(prs: Deck, f: dict[str, str]) -> None:  # noqa: ARG001 - общая подпись слайдов
    s = new_slide(prs, 4, "Рабочее место в браузере")
    top, h = 1.95, 3.3
    w_home = h * 16 / 9
    picture(s, "console.png", (LEFT, top, w_home, h), Crop())
    picture(s, "workspace.png", (LEFT + w_home + 0.3, top, WIDTH - w_home - 0.3, h), WORK_CROP)
    items = (
        "Главная: улица пилота или свой DXF/DWG и реестр прогонов на одном экране.",
        "Клик по кроне - замеры до сетей и пункт нормы в правой панели. Вкладки «Посадки», "
        "«Эффект», «Легенда».",
        "«Править»: перенос проверяется теми же нормами на лету, пересборка переписывает DXF.",
        "Отчёты, ведомость Excel и записка DOCX открываются окном поверх плана.",
    )
    gap = 0.3
    w = (WIDTH - 3 * gap) / 4
    for i, item in enumerate(items):
        x = LEFT + i * (w + gap)
        circle_number(s, x, 5.5, 0.42, str(i + 1), size=13)
        text(s, (x + 0.55, 5.5, w - 0.55, 1.3), [para(item, Style(MEDIUM, 11.5, INK))], line=1.2)
    notes(
        s,
        "Интерфейс второстепенен по ТЗ, но правка желательна: перенос, удаление, проверка точки "
        "на лету, пересборка DXF. Каталог моделей растений - отдельная страница.",
    )


# --------------------------------------------------------------------------- 05 уникальность


def s_unique(prs: Deck, f: dict[str, str]) -> None:
    s = new_slide(prs, 5, "Чем green отличается", bg="dark")
    cards = (
        (
            f"{f['rules_total']} правил · {f['acts']} актов",
            "Прозрачное объяснение",
            "У каждой посадки и каждого отказа замер, правило, акт, пункт и дословная цитата. "
            "Текст собирается из трассы проверок, без языковой модели.",
        ),
        (
            "2 редакции СП 42",
            "Нормы - данные, а не код",
            "Новая редакция акта - правка YAML без программиста и пересборки образа. "
            "СП 42.13330 2016 и 2026 уже в своде, выбор параметром.",
        ),
        (
            f"{f['violations']} нарушений",
            "Проверка после генерации",
            "Готовый план заново проверяется всеми правилами. Та же проверка находит нарушения в "
            "чужих планах: 391 из 554 посадок на Берзарина.",
        ),
        (
            f"{f['entities_changed']} из {f['entities']}",
            "Исходный чертёж не трогаем",
            "План пишется только на слои GREEN_*. После записи каждый объект исходника "
            "сверяется по отпечатку blake2b.",
        ),
        (
            f"{f['species']} видов",
            "Вид под место",
            "Задача MILP подбирает породу по условиям места: соль у проезда, тень домов, сети "
            "рядом, квоты 369-ПП на разнообразие пород.",
        ),
        (
            "1 план",
            "DXF, 3D и фото из одного плана",
            "Проектировщик получает чертёж, согласующий - объяснения, жители - вид улицы. "
            "Всё собрано из одного проверенного плана.",
        ),
    )
    gap_x, gap_y = 0.28, 0.24
    w = (WIDTH - 2 * gap_x) / 3
    h = 2.25
    for i, (metric, head, body) in enumerate(cards):
        x = LEFT + (i % 3) * (w + gap_x)
        y = 1.95 + (i // 3) * (h + gap_y)
        c = rect(s, (x, y, w, h), WHITE, line=(WHITE, 1.0), radius=0.16)
        alpha(c, 9, 26)
        text(s, (x + 0.28, y + 0.22, w - 0.56, 0.5), [para(metric, Style(X_BOLD, 19, ROSE))])
        text(s, (x + 0.28, y + 0.7, w - 0.56, 0.3), [para(head, Style(SEMI, 13, WHITE))], line=1.0)
        text(
            s,
            (x + 0.28, y + 1.1, w - 0.56, 1.05),
            [para(body, Style(MEDIUM, 10.5, MIST))],
            line=1.2,
        )
    foot(
        s,
        "У заказчика попыток автоматизации не было: green первым ставит посадки по сетям "
        "и доказывает каждую норму.",
        dark=True,
        y=6.75,
    )
    notes(
        s,
        "Уникальность: объяснимость без LLM, нормы как данные, независимая проверка и "
        "нормоконтроль чужих планов, неприкосновенность исходника, вид под место, единый план "
        "для DXF, 3D и фото.",
    )


# --------------------------------------------------------------------------- 06 развитие


def s_roadmap(prs: Deck, f: dict[str, str]) -> None:  # noqa: ARG001 - общая подпись слайдов
    s = new_slide(prs, 6, "Как сервис будет расти")
    stages = (
        (
            "Пилот в ДПиООС",
            (
                "образ Docker на МосТех.ОС: расчёт без сети и GPU",
                "проектировщики получают черновик улицы с объяснениями",
                "нормоконтроль проектов подрядчиков: green audit",
            ),
        ),
        (
            "Шире нормы и данные",
            (
                "водоохранные зоны и объекты наследия: нужен сверенный пункт о посадке",
                "охранные зоны по линиям сетей чертежа",
                "почвы и корнеобитаемый объём по данным изысканий",
            ),
        ),
        (
            "Масштаб",
            (
                "площадки, дорожки и малые формы скверов",
                "интеграция с системами департамента по OpenAPI",
                "сервер департамента на Kubernetes",
                "другие регионы: свой свод норм без правки кода",
            ),
        ),
    )
    gap = 0.4
    w = (WIDTH - 2 * gap) / 3
    for i, (head, items) in enumerate(stages):
        x = LEFT + i * (w + gap)
        circle_number(s, x, 2.0, 0.6, str(i + 1), size=17)
        text(
            s,
            (x + 0.8, 2.0, w - 0.8, 0.6),
            [para(f"Этап {i + 1}", Style(SEMI, 11, PINK)), para(head, Style(X_BOLD, 15, NIGHT))],
            anchor=MSO_ANCHOR.MIDDLE,
            line=1.05,
        )
        card = rect(s, (x, 2.9, w, 3.2), WHITE, line=(BLUSH, 1.25), radius=0.16)
        del card
        text(
            s,
            (x + 0.28, 3.15, w - 0.56, 2.8),
            [para(item, Style(MEDIUM, 12, BODY)) for item in items],
            line=1.2,
            after=9,
            bullet=PINK,
        )
    text(
        s,
        (LEFT, 6.4, WIDTH, 0.4),
        [
            [
                ("green", Style(X_BOLD, 16, PLUM)),
                ("   команда MISIS MOJARUNG · github.com/Mojarung/LCT_2026", Style(SEMI, 12, BODY)),
            ]
        ],
        anchor=MSO_ANCHOR.MIDDLE,
    )
    notes(
        s,
        "Путь внедрения по ТЗ: пилот на ресурсах ДПиООС, затем данные ГИС и изысканий, затем масштаб.",
    )


OUR_SLIDES: tuple[Callable[[Deck, dict[str, str]], None], ...] = (
    s_problem,
    s_solution,
    s_steps,
    s_explain,
    s_users,
    s_audit,
    s_visual,
    s_city,
    s_effect,
    s_deploy,
    s_speed,
    s_inputs,
    s_workspace,
    s_unique,
    s_roadmap,
)


def add_slides(prs: Deck) -> None:
    f = facts()
    for build in OUR_SLIDES:
        build(prs, f)
